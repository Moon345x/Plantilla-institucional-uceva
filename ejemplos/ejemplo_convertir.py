#!/usr/bin/env python3
"""Convierte un .pptx cualquiera a la plantilla UCEVA con una heurística simple.

    python ejemplos/ejemplo_convertir.py ORIGEN.pptx [DESTINO.pptx] [--etiqueta "Facultad de ..."]

Qué hace, diapositiva por diapositiva:
  1. Lee título, párrafos (con nivel de viñeta), tablas, imágenes y notas del orador.
  2. Decide el tipo de diapositiva de la plantilla (ver `elegir_tipo`).
  3. Reconstruye el contenido con la librería (NO copia las formas viejas).

Es un punto de partida: una IA debe revisar las decisiones marcadas con «DECIDE LA IA».
"""
import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "lib"))

from uceva_pptx import (ContenidoDemasiadoLargo, abrir_plantilla, cierre, diapositiva_contenido,  # noqa: E402
                        diapositiva_foto, dividir_tabla, guardar, imagen, portada, tabla, texto)

PALABRAS_CIERRE = ("gracias", "preguntas", "fin", "thank")


def leer(ruta: Path, carpeta_media: Path) -> list[dict]:
    """Extrae el contenido de cada diapositiva a un diccionario simple."""
    carpeta_media.mkdir(parents=True, exist_ok=True)
    origen = Presentation(str(ruta))
    res = []
    for n, s in enumerate(origen.slides, start=1):
        d = {"numero": n, "titulo": "", "subtitulo": "", "parrafos": [], "tablas": [], "imagenes": [], "notas": ""}
        for sh in s.shapes:
            if sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                d["titulo"] = sh.text_frame.text.strip()
            elif sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
                d["subtitulo"] = sh.text_frame.text.strip()
            elif sh.has_table if hasattr(sh, "has_table") else False:
                d["tablas"].append([[c.text.strip() for c in fila.cells] for fila in sh.table.rows])
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                destino = carpeta_media / f"slide{n:02d}_{sh.shape_id}.{sh.image.ext}"
                destino.write_bytes(sh.image.blob)
                d["imagenes"].append(destino)
            elif sh.has_text_frame and sh.text_frame.text.strip():
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        d["parrafos"].append((t, min(p.level, 1)))
        # DECIDE LA IA: si no hay marcador de título, aquí se toma el primer párrafo como título.
        if not d["titulo"] and d["parrafos"]:
            d["titulo"] = d["parrafos"].pop(0)[0]
        if s.has_notes_slide:
            d["notas"] = s.notes_slide.notes_text_frame.text.strip()
        res.append(d)
    return res


def elegir_tipo(d: dict, es_primera: bool, es_ultima: bool) -> str:
    """Heurística. DECIDE LA IA: revisa cada caso; por ejemplo, una diapositiva con 6 ideas
    paralelas queda mejor como cuadrícula de tarjetas, y una con una cifra, como numero_grande."""
    titulo = d["titulo"].lower()
    if es_primera:
        return "portada"
    if es_ultima and any(p in titulo for p in PALABRAS_CIERRE):
        return "cierre"
    if d["tablas"]:
        return "tabla"
    if d["imagenes"] and not d["parrafos"]:
        return "foto"
    if d["imagenes"]:
        return "texto_e_imagen"
    return "texto"


def convertir(origen: Path, destino: Path, etiqueta: str | None) -> None:
    media = destino.parent / (destino.stem + "_media")
    datos = leer(origen, media)
    prs = abrir_plantilla()
    for i, d in enumerate(datos):
        tipo = elegir_tipo(d, i == 0, i == len(datos) - 1)
        notas = d["notas"] or None
        if tipo == "portada":
            sub = d["subtitulo"] or (d["parrafos"][0][0] if d["parrafos"] else None)
            portada(prs, d["titulo"], sub, notas=notas)
        elif tipo == "cierre":
            sub = d["subtitulo"] or (d["parrafos"][0][0] if d["parrafos"] else None)
            cierre(prs, "¡Gracias!" if "gracias" in d["titulo"].lower() else d["titulo"], sub, notas=notas)
        elif tipo == "tabla":
            filas = d["tablas"][0]
            # Tabla larga: se parte en varias diapositivas repitiendo el encabezado.
            trozos = [filas]
            for max_filas in (len(filas), 8, 6, 4):
                trozos = dividir_tabla(filas, max_filas)
                try:
                    for k, trozo in enumerate(trozos):
                        t = d["titulo"] + (" (cont.)" if k else "")
                        s, area = diapositiva_contenido(prs, etiqueta, t, notas=notas)
                        tabla(s, area, trozo)
                    break
                except ContenidoDemasiadoLargo:
                    # deshacer las diapositivas creadas en este intento
                    for _ in range(k + 1):
                        _borrar_ultima(prs)
        elif tipo == "foto":
            diapositiva_foto(prs, d["imagenes"][0], d["titulo"], notas=notas)
        elif tipo == "texto_e_imagen":
            s, area = diapositiva_contenido(prs, etiqueta, d["titulo"], notas=notas)
            izq, der = area.columnas(2)
            texto(s, izq, d["parrafos"], vinetas=True, tam=18)
            imagen(s, d["imagenes"][0], der, borde=True)
        else:
            _texto_en_una_o_varias(prs, etiqueta, d, notas)
    guardar(prs, destino)
    print(f"Convertido: {origen.name} ({len(datos)} diapositivas) → {destino}")


def _texto_en_una_o_varias(prs, etiqueta, d, notas):
    """Pone los párrafos con viñetas; si no caben, los reparte en varias diapositivas."""
    pendientes = list(d["parrafos"]) or [("(sin texto)", 0)]
    k = 0
    while pendientes:
        t = d["titulo"] + (" (cont.)" if k else "")
        for n in range(len(pendientes), 0, -1):
            s, area = diapositiva_contenido(prs, etiqueta, t, notas=notas if k == 0 else None)
            try:
                texto(s, area, pendientes[:n], vinetas=len(pendientes) > 1, tam=20)
                pendientes = pendientes[n:]
                break
            except ContenidoDemasiadoLargo:
                _borrar_ultima(prs)
        k += 1


def _borrar_ultima(prs):
    lista = prs.slides._sldIdLst
    ultimo = lista[-1]
    lista.remove(ultimo)
    prs.part.drop_rel(ultimo.rId)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("origen", type=Path)
    ap.add_argument("destino", type=Path, nargs="?")
    ap.add_argument("--etiqueta", default=None, help="Texto de la etiqueta de facultad (opcional)")
    a = ap.parse_args()
    destino = a.destino or RAIZ / "salida" / f"{a.origen.stem}_uceva.pptx"
    convertir(a.origen, destino, a.etiqueta)


if __name__ == "__main__":
    main()
