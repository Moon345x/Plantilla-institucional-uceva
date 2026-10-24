#!/usr/bin/env python3
"""Genera salida/prueba_origen.pptx: una presentación con OTRO estilo (la plantilla por defecto
de python-pptx, 4:3) para probar ejemplo_convertir.py. Trae título, viñetas, tabla, imagen y notas.
Datos inventados."""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

RAIZ = Path(__file__).resolve().parent.parent
FOTO = RAIZ / "assets" / "img" / "foto_ejemplo_campus.jpeg"

prs = Presentation()   # plantilla blanca de python-pptx, 4:3

s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Introducción a las bases de datos"
s.placeholders[1].text = "Sesión 1 · Curso de ejemplo"
s.notes_slide.notes_text_frame.text = "Presentarse y explicar la dinámica del curso."

s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "¿Qué es una base de datos?"
tf = s.placeholders[1].text_frame
tf.text = "Un conjunto organizado de datos"
for t, nivel in [("Se guarda de forma persistente", 1), ("Se consulta con un lenguaje (SQL)", 0),
                 ("La administra un motor (SGBD)", 0), ("Ejemplos: PostgreSQL, MySQL, SQLite", 1)]:
    p = tf.add_paragraph()
    p.text, p.level = t, nivel
s.notes_slide.notes_text_frame.text = "Pedir ejemplos de bases de datos que usan a diario."

s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Tipos de datos más comunes"
filas = [["Tipo", "Ejemplo", "Uso"], ["INTEGER", "42", "Contadores, edades"], ["VARCHAR", "'Ana'", "Nombres, textos cortos"],
         ["DATE", "2026-03-15", "Fechas"], ["BOOLEAN", "TRUE", "Sí / no"], ["DECIMAL", "19.99", "Dinero"]]
tbl = s.shapes.add_table(len(filas), 3, Inches(0.5), Inches(1.6), Inches(9), Inches(3)).table
for i, fila in enumerate(filas):
    for j, v in enumerate(fila):
        tbl.cell(i, j).text = v
s.notes_slide.notes_text_frame.text = "Insistir en que DECIMAL es mejor que FLOAT para dinero."

s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Laboratorio de práctica"
s.shapes.add_picture(str(FOTO), Inches(2), Inches(1.7), height=Inches(4.5))
s.notes_slide.notes_text_frame.text = "Las prácticas son en la sala 3 (dato inventado)."

s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Para la próxima clase"
s.placeholders[1].text = "Instalar SQLite y abrir la base de ejemplo del curso."

s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Gracias"
s.placeholders[1].text = "¿Preguntas?"

salida = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "salida" / "prueba_origen.pptx"
salida.parent.mkdir(parents=True, exist_ok=True)
prs.save(salida)
print(f"Creado: {salida}")
