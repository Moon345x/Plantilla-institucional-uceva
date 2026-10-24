"""uceva_pptx: crear diapositivas con la plantilla institucional UCEVA 55 años sin tocar XML.

Uso mínimo:

    from uceva_pptx import *
    prs = abrir_plantilla()
    portada(prs, "Mi presentación", "Curso de ejemplo", autor="Equipo docente")
    s, area = diapositiva_contenido(prs, "Facultad de Ingeniería", "Objetivo de la sesión")
    texto(s, area, ["Primer punto", "Segundo punto"], vinetas=True)
    cierre(prs)
    guardar(prs, "salida/mi_presentacion.pptx")

Todas las medidas y colores salen de spec/tokens.json (generado desde el .pptx).
Unidades: EMU (914400 por pulgada, 12700 por punto). Las funciones que dibujan
devuelven el Area que ocuparon, para poder seguir debajo con `area.debajo(usada)`.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import os
import re
import weakref
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

__all__ = [
    "Area", "ContenidoDemasiadoLargo", "TOKENS", "COLORES",
    "abrir_plantilla", "guardar",
    "bienvenida", "portada", "diapositiva_contenido", "diapositiva_barra_lateral",
    "diapositiva_en_blanco", "diapositiva_foto", "diapositiva_lista_de_bloques", "cierre",
    "copiar_muestra", "reemplazar_texto",
    "texto", "tarjeta", "fila_de_tarjetas", "cuadricula", "tabla", "dividir_tabla",
    "bloque_codigo", "imagen", "numero_grande", "fila_de_numeros",
    "medir_alto_texto", "color_texto_sobre",
]

RAIZ = Path(__file__).resolve().parent.parent
RUTA_TOKENS = RAIZ / "spec" / "tokens.json"
PLANTILLA_POR_DEFECTO = RAIZ / "plantilla" / "PLANTILLA-PRESENTACION-INSTITUCIONAL-55-anos_VF.pptx"

if not RUTA_TOKENS.exists():
    raise FileNotFoundError(f"Falta {RUTA_TOKENS}. Ejecuta: python scripts/extraer_plantilla.py")
TOKENS = json.loads(RUTA_TOKENS.read_text(encoding="utf-8"))

COLORES = {k: v["hex"] for k, v in TOKENS["colores"].items()}
COLORES.update({k: v["hex"] for k, v in TOKENS["colores_derivados"].items()})
_EL = TOKENS["elementos"]
_ZONAS = TOKENS["zonas_protegidas"]
_TAM = TOKENS["tamanos_pt"]

ANCHO_HOJA = TOKENS["hoja"]["ancho_emu"]
ALTO_HOJA = TOKENS["hoja"]["alto_emu"]
EMU_PT = TOKENS["hoja"]["emu_por_punto"]
EMU_IN = TOKENS["hoja"]["emu_por_pulgada"]

# Margen izquierdo donde empieza el texto en la plantilla: caja en x=375167 + margen interno 0,1".
X_TEXTO = _EL["titulo_contenido"]["posicion"]["emu"]["x"] + 91440
MARGEN_INF = TOKENS["margenes"]["inferior_minimo_emu"]
SEPARACION = int(0.2 * EMU_IN)          # espacio estándar entre elementos
LIMITE_INF = ALTO_HOJA - MARGEN_INF

FUENTE = TOKENS["fuentes"]["texto"]
FUENTE_PORTADA = TOKENS["fuentes"]["titulo_portada"]
FUENTE_CODIGO = TOKENS["fuentes"]["codigo"]

COLORES_BLOQUES = ["morado", "rosado", "amarillo", "cian"]
TEXTOS_DE_EJEMPLO = TOKENS["textos_de_ejemplo"]

_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


class ContenidoDemasiadoLargo(ValueError):
    """El contenido no cabe en el espacio disponible: divídelo en más diapositivas."""


# ---------------------------------------------------------------------------
# Área de trabajo
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Area:
    """Rectángulo en EMU. `diapositiva_contenido` devuelve el área libre para el contenido."""
    x: int
    y: int
    ancho: int
    alto: int

    @property
    def abajo(self) -> int:
        return self.y + self.alto

    @property
    def derecha(self) -> int:
        return self.x + self.ancho

    def debajo(self, usada: "Area | int", separacion: int = SEPARACION) -> "Area":
        """Lo que queda de esta área por debajo de `usada` (un Area o un alto en EMU)."""
        fin = usada.abajo if isinstance(usada, Area) else self.y + int(usada)
        nuevo_y = fin + separacion
        return Area(self.x, nuevo_y, self.ancho, max(0, self.abajo - nuevo_y))

    def columnas(self, n: int, separacion: int = SEPARACION, proporciones: list[float] | None = None) -> list["Area"]:
        """Divide el área en `n` columnas (iguales o según `proporciones`)."""
        proporciones = proporciones or [1] * n
        util = self.ancho - separacion * (len(proporciones) - 1)
        total = sum(proporciones)
        res, x = [], self.x
        for p in proporciones:
            a = int(util * p / total)
            res.append(Area(x, self.y, a, self.alto))
            x += a + separacion
        return res

    def filas(self, n: int, separacion: int = SEPARACION) -> list["Area"]:
        alto = (self.alto - separacion * (n - 1)) // n
        return [Area(self.x, self.y + i * (alto + separacion), self.ancho, alto) for i in range(n)]

    def con_alto(self, alto: int) -> "Area":
        return Area(self.x, self.y, self.ancho, int(alto))

    def pulgadas(self) -> tuple[float, float, float, float]:
        return tuple(round(v / EMU_IN, 2) for v in (self.x, self.y, self.ancho, self.alto))


# ---------------------------------------------------------------------------
# Medición de texto (para que nada se desborde)
# ---------------------------------------------------------------------------

_ARCHIVOS_FUENTE = {
    ("Arial", False): ["arial.ttf", "Arial.ttf", "LiberationSans-Regular.ttf", "Arimo-Regular.ttf"],
    ("Arial", True): ["arialbd.ttf", "Arial Bold.ttf", "LiberationSans-Bold.ttf", "Arimo-Bold.ttf"],
    ("Arial Black", False): ["ariblk.ttf", "Arial Black.ttf"],
    ("Courier New", False): ["cour.ttf", "Courier New.ttf", "LiberationMono-Regular.ttf", "Cousine-Regular.ttf"],
    ("Courier New", True): ["courbd.ttf", "Courier New Bold.ttf", "LiberationMono-Bold.ttf", "Cousine-Bold.ttf"],
}
# Ancho medio de carácter (en "em") si no hay ninguna fuente instalada para medir.
_ANCHO_MEDIO = {"Arial": 0.52, "Arial Black": 0.70, "Courier New": 0.60}


@lru_cache(maxsize=1)
def _indice_fuentes() -> dict[str, str]:
    dirs = ["/usr/share/fonts", "/usr/local/share/fonts", os.path.expanduser("~/.local/share/fonts"),
            os.path.expanduser("~/.fonts"), "C:/Windows/Fonts", "/Library/Fonts",
            "/System/Library/Fonts/Supplemental", os.path.expanduser("~/Library/Fonts")]
    idx = {}
    for d in dirs:
        for raiz, _, archivos in os.walk(d):
            for a in archivos:
                idx.setdefault(a.lower(), os.path.join(raiz, a))
    return idx


@lru_cache(maxsize=None)
def _fuente_pil(familia: str, negrita: bool):
    idx = _indice_fuentes()
    for nombre in _ARCHIVOS_FUENTE.get((familia, negrita), []):
        if nombre.lower() in idx:
            return ImageFont.truetype(idx[nombre.lower()], 100), 1.0
    if familia == "Arial Black":   # sin Arial Black: Arial negrita ensanchada (Arial Black es ~10-20 % más ancha)
        f, _ = _fuente_pil("Arial", True)
        return f, 1.15
    return None, 1.0


def _ancho_pt(texto: str, pt: float, familia: str = FUENTE, negrita: bool = False) -> float:
    f, factor = _fuente_pil(familia, negrita)
    if f is None:
        return len(texto) * pt * _ANCHO_MEDIO.get(familia, 0.55) * (1.08 if negrita else 1)
    return f.getlength(texto) * pt / 100 * factor


def _contar_lineas(texto: str, pt: float, ancho_emu: int, familia: str = FUENTE, negrita: bool = False) -> int:
    """Líneas que ocupa `texto` en una caja de `ancho_emu` (corte por palabras, como PowerPoint)."""
    disponible = ancho_emu / EMU_PT * 0.97        # 3 % de holgura por diferencias de motor
    total = 0
    for parrafo in texto.split("\n"):
        palabras = parrafo.split(" ")
        lineas, actual = 1, ""
        for p in palabras:
            prueba = p if not actual else actual + " " + p
            if _ancho_pt(prueba, pt, familia, negrita) <= disponible:
                actual = prueba
                continue
            if actual:
                lineas += 1
            actual = p
            while _ancho_pt(actual, pt, familia, negrita) > disponible and len(actual) > 1:  # palabra más ancha que la caja
                corte = len(actual)
                while corte > 1 and _ancho_pt(actual[:corte], pt, familia, negrita) > disponible:
                    corte -= 1
                actual = actual[corte:]
                lineas += 1
        total += lineas
    return total


INTERLINEADO = 1.2   # alto de línea / tamaño de letra (Arial en PowerPoint ≈ 1,15; 1,2 deja margen)


def medir_alto_texto(parrafos: list[str] | str, pt: float, ancho_emu: int, negrita: bool = False,
                     familia: str = FUENTE, espacio_parrafos_pt: float = 0, sangria_emu: int = 0) -> int:
    """Alto en EMU (sin márgenes internos) que necesitan los párrafos dentro de `ancho_emu`."""
    if isinstance(parrafos, str):
        parrafos = [parrafos]
    lineas = sum(_contar_lineas(_sin_marcas(p), pt, ancho_emu - sangria_emu, familia, negrita or "**" in p) for p in parrafos)
    extra = espacio_parrafos_pt * max(0, len(parrafos) - 1)
    return int((lineas * pt * INTERLINEADO + extra) * EMU_PT)


def _sin_marcas(t: str) -> str:
    return t.replace("**", "")


# ---------------------------------------------------------------------------
# Colores
# ---------------------------------------------------------------------------

def _hex(color: str) -> str:
    """Acepta un nombre de tokens.json ('morado') o un hex ('6000AF'). Sin '#'."""
    if color in COLORES:
        return COLORES[color]
    c = color.lstrip("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}", c):
        raise ValueError(f"Color no válido: {color!r}. Usa un nombre de tokens.json: {sorted(COLORES)}")
    return c


def _luminancia(hexa: str) -> float:
    def canal(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (int(hexa[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * canal(r) + 0.7152 * canal(g) + 0.0722 * canal(b)


def _contraste(a: str, b: str) -> float:
    la, lb = sorted((_luminancia(_hex(a)), _luminancia(_hex(b))), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def color_texto_sobre(fondo: str) -> str:
    """'blanco' o 'gris_texto', el que más contraste tenga sobre `fondo`."""
    return "blanco" if _contraste("blanco", fondo) >= _contraste("gris_texto", fondo) else "gris_texto"


# ---------------------------------------------------------------------------
# Ayudas de bajo nivel (formas, texto, XML)
# ---------------------------------------------------------------------------

def _rgb(color: str) -> RGBColor:
    return RGBColor.from_string(_hex(color))


def _poner_fuente(run, familia: str = FUENTE):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", familia)


def _fin_parrafo(p, tam, familia=FUENTE):
    """Marca de fin de párrafo con el mismo tamaño y fuente: así una línea vacía no hereda 18 pt de Calibri."""
    end = p._p.get_or_add_endParaRPr()
    end.set("sz", str(int(tam * 100)))
    for tag in ("a:latin", "a:cs"):
        etree.SubElement(end, qn(tag)).set("typeface", familia)


def _sombra_texto(run):
    """Sombra suave del texto de las etiquetas (copiada de la plantilla)."""
    xml = _EL["etiqueta_facultad"]["sombra_texto_xml"].replace("<a:outerShdw ", f'<a:outerShdw xmlns:a="{_NS_A}" ', 1)
    rPr = run._r.get_or_add_rPr()
    efecto = etree.SubElement(rPr, qn("a:effectLst"))
    efecto.append(etree.fromstring(xml))
    # effectLst debe ir antes de latin/cs según el esquema
    for tag in ("a:latin", "a:cs"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
            rPr.append(el)


def _escribir(tf, parrafos, tam, color, negrita=False, alineacion="izquierda", familia=FUENTE,
              vinetas=False, espacio_parrafos_pt=0, sombra=False):
    """Escribe párrafos en un text_frame. `**texto**` va en negrita. Un párrafo puede ser (texto, nivel)."""
    alin = {"izquierda": PP_ALIGN.LEFT, "centro": PP_ALIGN.CENTER, "derecha": PP_ALIGN.RIGHT,
            "justificado": PP_ALIGN.JUSTIFY}[alineacion]
    tf.clear()
    for i, par in enumerate(parrafos):
        nivel = 0
        if isinstance(par, tuple):
            par, nivel = par
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alin
        if i > 0 and espacio_parrafos_pt:
            p.space_before = Pt(espacio_parrafos_pt)
        if vinetas:
            pPr = p._p.get_or_add_pPr()
            sangria = int(0.25 * EMU_IN) * (nivel + 1)
            pPr.set("marL", str(sangria))
            pPr.set("indent", str(-int(0.25 * EMU_IN)))
            bf = etree.SubElement(pPr, qn("a:buFont"))
            bf.set("typeface", FUENTE)
            bc = etree.SubElement(pPr, qn("a:buChar"))
            bc.set("char", "•" if nivel == 0 else "–")
        trozos = re.split(r"(\*\*.+?\*\*)", str(par))
        for t in trozos:
            if not t:
                continue
            r = p.add_run()
            en_negrita = t.startswith("**") and t.endswith("**")
            r.text = t[2:-2] if en_negrita else t
            r.font.size = Pt(tam)
            r.font.bold = negrita or en_negrita
            r.font.color.rgb = _rgb(color)
            if sombra:
                _sombra_texto(r)
            _poner_fuente(r, familia)
        _fin_parrafo(p, tam, familia)


def _caja_texto(slide, x, y, ancho, alto, nombre, margen_x=0, margen_y=0, anclaje="arriba", ajustar=False):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(ancho), Emu(alto))
    tb.name = nombre
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(margen_x)
    tf.margin_top = tf.margin_bottom = Emu(margen_y)
    tf.vertical_anchor = {"arriba": MSO_ANCHOR.TOP, "centro": MSO_ANCHOR.MIDDLE, "abajo": MSO_ANCHOR.BOTTOM}[anclaje]
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for hijo in list(bodyPr):
        bodyPr.remove(hijo)
    if ajustar:
        etree.SubElement(bodyPr, qn("a:spAutoFit"))
    return tb


def _forma(slide, tipo, x, y, ancho, alto, nombre, relleno=None, linea=None, ancho_linea_pt=1.0,
           guion=False, opacidad=None):
    geom = {"rect": MSO_SHAPE.RECTANGLE, "redondeado": MSO_SHAPE.ROUNDED_RECTANGLE}[tipo]
    sh = slide.shapes.add_shape(geom, Emu(x), Emu(y), Emu(ancho), Emu(alto))
    sh.name = nombre
    if relleno:
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(relleno)
        if opacidad is not None:   # transparencia con <a:alpha>, nunca con hex de 8 dígitos
            clr = sh.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(opacidad * 100000)))
    else:
        sh.fill.background()
    if linea:
        sh.line.color.rgb = _rgb(linea)
        sh.line.width = Pt(ancho_linea_pt)
        if guion:
            ln = sh.line._get_or_add_ln()
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False      # sin sombra heredada del estilo
    tf = sh.text_frame
    tf.word_wrap = True
    return sh


def _quitar_marcadores(slide):
    """Elimina los marcadores (placeholders) que python-pptx copia del diseño. Si quedan vacíos,
    PowerPoint muestra 'Haga clic para agregar título' y el validador los marca como error."""
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)


def _poner_notas(slide, notas):
    if notas:
        slide.notes_slide.notes_text_frame.text = notas


def _fondo_imagen(slide, nombre_imagen):
    parte = _estado(slide)["imagenes"][nombre_imagen]
    rId = slide.part.relate_to(parte, RT.IMAGE)
    cSld = slide._element.find(qn("p:cSld"))
    viejo = cSld.find(qn("p:bg"))
    if viejo is not None:
        cSld.remove(viejo)
    bg = etree.fromstring(
        f'<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="{_NS_A}" xmlns:r="{_NS_R}">'
        f'<p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rId}"/><a:srcRect/>'
        f'<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    cSld.insert(0, bg)


def _imagen_de_plantilla(slide, nombre_imagen, x, y, ancho, alto, nombre):
    parte = _estado(slide)["imagenes"][nombre_imagen]
    # python-pptx reconoce la imagen por su SHA1 y reutiliza la parte existente (no la duplica)
    pic = slide.shapes.add_picture(io.BytesIO(parte.blob), Emu(x), Emu(y), Emu(ancho), Emu(alto))
    pic.name = nombre
    return pic


# Estado de cada presentación abierta con abrir_plantilla (muestras, cierres, imágenes, layouts).
_ESTADOS: "weakref.WeakKeyDictionary" = weakref.WeakKeyDictionary()


def _estado(obj) -> dict:
    """Estado de la presentación a la que pertenece `obj` (presentación o diapositiva)."""
    parte = obj.part if hasattr(obj, "slides") else obj.part.package.presentation_part
    try:
        return _ESTADOS[parte]
    except KeyError:
        raise RuntimeError("Abre la presentación con abrir_plantilla(), no con Presentation().") from None


def _nueva(prs, layout="contenido"):
    s = prs.slides.add_slide(_estado(prs)["layouts"][layout])
    _quitar_marcadores(s)
    return s


# ---------------------------------------------------------------------------
# Abrir y guardar
# ---------------------------------------------------------------------------

def abrir_plantilla(ruta: str | Path | None = None) -> Presentation:
    """Abre la plantilla (sin modificar el archivo original) y prepara lo que usa la librería."""
    ruta = Path(ruta) if ruta else PLANTILLA_POR_DEFECTO
    prs = Presentation(str(ruta))
    por_sha = {v["sha1"]: k for k, v in TOKENS["imagenes"].items()}
    imagenes = {}
    for parte in prs.part.package.iter_parts():
        if parte.partname.startswith("/ppt/media/"):
            nombre = por_sha.get(hashlib.sha1(parte.blob).hexdigest())
            if nombre:
                imagenes[nombre] = parte
    layouts = {l.name: l for l in prs.slide_layouts}
    _ESTADOS[prs.part] = {
        "muestras": [s.slide_id for s in prs.slides],
        "cierres": [],
        "imagenes": imagenes,
        "layouts": {"contenido": layouts.get("Diapositiva de título", prs.slide_layouts[0]),
                    "en_blanco": layouts.get("En blanco", prs.slide_layouts[6])},
    }
    faltan = {"fondo_portada", "fondo_bienvenidos", "fondo_barra_lateral", "logo_55_uceva_blanco", "fondo_cierre"} - set(imagenes)
    if faltan:
        raise RuntimeError(f"La plantilla no coincide con spec/tokens.json (faltan {faltan}). "
                           "Vuelve a ejecutar scripts/extraer_plantilla.py.")
    return prs


def guardar(prs, ruta: str | Path, conservar_muestras: bool = False) -> Path:
    """Elimina las 11 diapositivas de muestra (soltando su relación para no dejar partes
    huérfanas), deja el cierre al final y guarda."""
    est = _estado(prs)
    lista = prs.slides._sldIdLst
    if not conservar_muestras:
        for sldId in list(lista):
            if sldId.id in est["muestras"]:
                rId = sldId.rId
                lista.remove(sldId)
                prs.part.drop_rel(rId)
    for sldId in list(lista):
        if sldId.id in est["cierres"]:
            lista.remove(sldId)
            lista.append(sldId)
    for s in prs.slides:
        for ph in list(s.placeholders):
            if not ph.has_text_frame or not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)
    prs.part.rename_slide_parts([s.rId for s in lista])
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(ruta))
    return ruta


# ---------------------------------------------------------------------------
# Tipos de diapositiva
# ---------------------------------------------------------------------------

def bienvenida(prs, notas: str | None = None):
    """Muestra 1: solo la imagen 'Bienvenidos'."""
    s = _nueva(prs)
    _fondo_imagen(s, "fondo_bienvenidos")
    _poner_notas(s, notas)
    return s


def portada(prs, titulo: str, subtitulo: str | None = None, autor: str | None = None,
            fecha: str | None = None, notas: str | None = None):
    """Muestra 2: título en la caja morada (Arial Black), marco punteado cian y subtítulo.
    Si el título es largo baja la letra (32 → 24 pt) y crece hacia abajo; el subtítulo
    siempre se coloca debajo de la caja, nunca encima."""
    s = _nueva(prs)
    _fondo_imagen(s, "fondo_portada")
    caja_t = _EL["titulo_portada"]["posicion"]["emu"]
    marco = _EL["marco_portada"]["posicion"]["emu"]
    sub = _EL["subtitulo_portada"]["posicion"]["emu"]
    margen_caja = int(0.15 * EMU_IN)
    ancho_util = caja_t["ancho"] - 2 * margen_caja
    tam = None
    for t in range(int(_TAM["titulo_portada"]), int(_TAM["titulo_portada_minimo"]) - 1, -2):
        if _contar_lineas(titulo, t, ancho_util, FUENTE_PORTADA) <= (1 if t == _TAM["titulo_portada"] else 3):
            tam = t
            break
    if tam is None:
        raise ContenidoDemasiadoLargo(f"Título de portada demasiado largo ({len(titulo)} caracteres). "
                                      "Acórtalo (máx. ~60) o pasa parte al subtítulo.")
    alto_t = max(caja_t["alto"], medir_alto_texto(titulo, tam, ancho_util, familia=FUENTE_PORTADA) + 2 * int(0.08 * EMU_IN))

    # Subtítulo y autor: ancho desde la x del subtítulo de la plantilla hasta el margen derecho.
    x_sub = sub["x"] + 91440
    ancho_sub = ANCHO_HOJA - X_TEXTO - x_sub
    y = caja_t["y"] + alto_t + int(0.1 * EMU_IN)
    bloques = []
    if subtitulo:
        tam_sub = int(_EL["subtitulo_portada"]["texto"]["tam_pt"])
        while tam_sub > 18 and _contar_lineas(subtitulo, tam_sub, ancho_sub, negrita=True) > 1:
            tam_sub -= 2
        bloques.append((subtitulo, tam_sub, True, "morado"))
    for extra in (autor, fecha):
        if extra:
            bloques.append((extra, 16, False, "gris_texto"))
    alturas = [medir_alto_texto(t, tm, ancho_sub, negrita=n) for t, tm, n, _ in bloques]
    fin = y + sum(alturas) + int(0.08 * EMU_IN) * max(0, len(bloques) - 1)
    alto_marco = max(marco["alto"], fin + int(0.25 * EMU_IN) - marco["y"])
    if marco["y"] + alto_marco > _ZONAS["arte_bienvenidos_portada"]["y"] + _ZONAS["arte_bienvenidos_portada"]["alto"]:
        raise ContenidoDemasiadoLargo("La portada no cabe: acorta el subtítulo o quita la fecha.")

    # Orden de capas: marco (atrás) → caja del título → textos.
    _forma(s, "redondeado", marco["x"], marco["y"], marco["ancho"], alto_marco, "uceva:marco_portada",
           linea="cian", ancho_linea_pt=_EL["marco_portada"]["linea"]["ancho_pt"], guion=True)
    caja = _forma(s, "redondeado", caja_t["x"], caja_t["y"], caja_t["ancho"], alto_t, "uceva:titulo_portada", relleno="morado")
    tf = caja.text_frame
    tf.margin_left = tf.margin_right = Emu(margen_caja)
    tf.margin_top = tf.margin_bottom = Emu(int(0.04 * EMU_IN))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _escribir(tf, [titulo], tam, "blanco", alineacion="centro", familia=FUENTE_PORTADA, sombra=True)
    for (t, tm, n, c), h in zip(bloques, alturas):
        tb = _caja_texto(s, x_sub, y, ancho_sub, h, "uceva:subtitulo_portada" if n else "uceva:autor")
        _escribir(tb.text_frame, [t], tm, c, negrita=n)
        y += h + int(0.08 * EMU_IN)
    _poner_notas(s, notas)
    return s


def _etiqueta(s, texto_etq: str, color: str):
    """Etiqueta de facultad arriba a la derecha (rectángulo redondeado, 16 pt negrita)."""
    e = _EL["etiqueta_facultad"]["posicion"]["emu"]
    tam = int(_EL["etiqueta_facultad"]["texto"]["tam_pt"])
    derecha = e["x"] + e["ancho"]
    limite_izq = _ZONAS["logo_contenido"]["ancho"] + SEPARACION
    while True:
        necesario = int(_ancho_pt(texto_etq, tam, negrita=True) * EMU_PT / 0.97) + 2 * int(0.2 * EMU_IN)
        if necesario <= derecha - limite_izq or tam <= 12:
            break
        tam -= 1
    ancho = min(max(e["ancho"], necesario), derecha - limite_izq)
    if _contar_lineas(texto_etq, tam, ancho - 2 * int(0.1 * EMU_IN), negrita=True) > 1:
        raise ContenidoDemasiadoLargo(f"Etiqueta demasiado larga: {texto_etq!r}. Usa una forma corta.")
    f = _forma(s, "redondeado", derecha - ancho, e["y"], ancho, e["alto"], "uceva:etiqueta", relleno=color)
    tf = f.text_frame
    tf.margin_left = tf.margin_right = Emu(int(0.1 * EMU_IN))
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    c_txt = color_texto_sobre(color)
    _escribir(tf, [texto_etq], tam, c_txt, negrita=True, alineacion="centro", sombra=(c_txt == "blanco"))
    return f


def _titulo_y_subtitulo(s, x, y, ancho, titulo, subtitulo, color_titulo):
    """Título 24 pt negrita + subtítulo opcional 18 pt gris. Devuelve el y donde terminan."""
    tam_t = int(_EL["titulo_contenido"]["texto"]["tam_pt"])
    if _contar_lineas(titulo, tam_t, ancho, negrita=True) > 2:
        raise ContenidoDemasiadoLargo(f"Título demasiado largo para dos líneas: {titulo!r}")
    alto_t = medir_alto_texto(titulo, tam_t, ancho, negrita=True)
    tb = _caja_texto(s, x, y, ancho, alto_t, "uceva:titulo")
    _escribir(tb.text_frame, [titulo], tam_t, color_titulo, negrita=True)
    fin = y + alto_t
    if subtitulo:
        tam_s = _TAM["subtitulo"]
        if _contar_lineas(subtitulo, tam_s, ancho) > 2:
            raise ContenidoDemasiadoLargo(f"Subtítulo demasiado largo: {subtitulo!r}. Pásalo al cuerpo.")
        alto_s = medir_alto_texto(subtitulo, tam_s, ancho)
        sb = _caja_texto(s, x, fin + int(0.04 * EMU_IN), ancho, alto_s, "uceva:subtitulo")
        _escribir(sb.text_frame, [subtitulo], tam_s, "gris_texto")
        fin += int(0.04 * EMU_IN) + alto_s
    return fin


def diapositiva_contenido(prs, etiqueta: str | None, titulo: str, subtitulo: str | None = None,
                          color: str = "morado", notas: str | None = None):
    """Muestras 4-6: logo (del fondo) + etiqueta de color arriba a la derecha + título.
    `color`: 'morado' (6), 'rosado' (5) o 'amarillo' (4). El título va en morado, o en rosado
    si color='rosado' (el amarillo no se usa como texto sobre blanco).
    Devuelve (diapositiva, Area libre para el contenido)."""
    if color not in ("morado", "rosado", "amarillo", "cian"):
        raise ValueError("color debe ser 'morado', 'rosado', 'amarillo' o 'cian'")
    s = _nueva(prs)
    if etiqueta:
        _etiqueta(s, etiqueta, color)
    y0 = _EL["titulo_contenido"]["posicion"]["emu"]["y"]
    ancho = ANCHO_HOJA - 2 * X_TEXTO
    fin = _titulo_y_subtitulo(s, X_TEXTO, y0, ancho, titulo, subtitulo, "rosado" if color == "rosado" else "morado")
    _poner_notas(s, notas)
    y = fin + SEPARACION
    return s, Area(X_TEXTO, y, ancho, LIMITE_INF - y)


def diapositiva_barra_lateral(prs, titulo: str, subtitulo: str | None = None, color: str = "morado",
                              notas: str | None = None):
    """Muestras 7-8: barra de color a la izquierda con el logo blanco; título a la derecha.
    Devuelve (diapositiva, Area libre)."""
    if color not in ("morado", "rosado"):
        raise ValueError("color debe ser 'morado' o 'rosado' (el logo es blanco)")
    s = _nueva(prs)
    _fondo_imagen(s, "fondo_barra_lateral")
    b = _EL["barra_lateral"]["posicion"]["emu"]
    _forma(s, "rect", b["x"], b["y"], b["ancho"], b["alto"], "uceva:barra", relleno=color)
    lg = _EL["logo_blanco_barra"]["posicion"]["emu"]
    _imagen_de_plantilla(s, "logo_55_uceva_blanco", lg["x"], lg["y"], lg["ancho"], lg["alto"], "uceva:logo")
    t = _EL["texto_con_barra"]["posicion"]["emu"]
    x = t["x"] + 91440
    ancho = ANCHO_HOJA - X_TEXTO - x
    fin = _titulo_y_subtitulo(s, x, t["y"] + 45720, ancho, titulo, subtitulo, color)
    _poner_notas(s, notas)
    y = fin + SEPARACION
    return s, Area(x, y, ancho, LIMITE_INF - y)


def diapositiva_en_blanco(prs, notas: str | None = None):
    """Muestra 3: solo el fondo con el logo. Devuelve (diapositiva, Area libre bajo el logo)."""
    s = _nueva(prs, "en_blanco")
    _poner_notas(s, notas)
    y = _ZONAS["logo_contenido"]["alto"] + SEPARACION
    return s, Area(X_TEXTO, y, ANCHO_HOJA - 2 * X_TEXTO, LIMITE_INF - y)


def diapositiva_foto(prs, ruta_imagen: str | Path, pie: str, texto: str | None = None,
                     notas: str | None = None):
    """Muestra 9: foto centrada con borde blanco y sombra, etiqueta morada sobre el borde
    inferior (`pie`) y un texto corto opcional debajo."""
    s = _nueva(prs)
    f = _EL["foto"]["posicion"]["emu"]
    etq = _EL["pie_foto_etiqueta"]["posicion"]["emu"]
    tam_pie = int(_EL["pie_foto_etiqueta"]["texto"]["tam_pt"])
    ancho_txt = int(7.5 * EMU_IN)
    alto_txt = medir_alto_texto(texto, 14, ancho_txt) if texto else 0
    sobresale = etq["y"] + etq["alto"] - (f["y"] + f["alto"])       # la etiqueta baja del borde de la foto
    reserva = sobresale + (alto_txt + int(0.05 * EMU_IN) if texto else 0)
    max_alto = min(f["alto"], LIMITE_INF - reserva - f["y"])
    with Image.open(ruta_imagen) as im:
        prop = im.width / im.height
    ancho, alto = f["ancho"], int(f["ancho"] / prop)
    if alto > max_alto:
        alto, ancho = max_alto, int(max_alto * prop)
    x = (ANCHO_HOJA - ancho) // 2
    pic = s.shapes.add_picture(str(ruta_imagen), Emu(x), Emu(f["y"]), Emu(ancho), Emu(alto))
    pic.name = "uceva:foto"
    _borde_y_sombra(pic)
    abajo = f["y"] + alto
    ancho_etq = min(max(etq["ancho"], int(_ancho_pt(pie, tam_pie, negrita=True) * EMU_PT / 0.97) + int(0.4 * EMU_IN)),
                    ANCHO_HOJA - 2 * X_TEXTO)
    if _contar_lineas(pie, tam_pie, ancho_etq - int(0.2 * EMU_IN), negrita=True) > 1:
        raise ContenidoDemasiadoLargo(f"Pie de foto demasiado largo: {pie!r}. Pon el detalle en `texto`.")
    y_etq = abajo - (etq["alto"] - sobresale)
    e = _forma(s, "redondeado", (ANCHO_HOJA - ancho_etq) // 2, y_etq, ancho_etq, etq["alto"], "uceva:pie_etiqueta", relleno="morado")
    e.text_frame.margin_top = e.text_frame.margin_bottom = Emu(0)
    e.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _escribir(e.text_frame, [pie], tam_pie, "blanco", negrita=True, alineacion="centro", sombra=True)
    if texto:
        tb = _caja_texto(s, (ANCHO_HOJA - ancho_txt) // 2, y_etq + etq["alto"] + int(0.05 * EMU_IN), ancho_txt, alto_txt, "uceva:pie")
        _escribir(tb.text_frame, [texto], 14, "gris_texto", alineacion="centro")
    _poner_notas(s, notas)
    return s


def _borde_y_sombra(pic):
    """Borde blanco de 6 pt y sombra suave, como la foto de la muestra 9."""
    spPr = pic._element.spPr
    ln = etree.SubElement(spPr, qn("a:ln"))
    ln.set("w", str(int(_EL["foto"]["linea"]["ancho_pt"] * EMU_PT)))
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", COLORES["blanco"])
    xml = _EL["foto"]["sombra_xml"].replace("<a:outerShdw ", f'<a:outerShdw xmlns:a="{_NS_A}" ', 1)
    ef = etree.SubElement(spPr, qn("a:effectLst"))
    ef.append(etree.fromstring(xml))


def diapositiva_lista_de_bloques(prs, bloques: list[tuple[str, str]], titulo: str | None = None,
                                 imagen_derecha: str | Path | None = None, notas: str | None = None):
    """Muestra 10: bloques apilados (etiqueta de color + texto corto). Hasta 4 en una columna;
    de 5 a 8 en dos columnas. Colores en orden: morado, rosado, amarillo, cian.
    `imagen_derecha` (opcional) ocupa la mitad derecha cuando hay una sola columna."""
    if not 1 <= len(bloques) <= 8:
        raise ContenidoDemasiadoLargo("Usa entre 1 y 8 bloques; divide el resto en otra diapositiva.")
    s = _nueva(prs)
    b0 = next(b for b in _EL["bloques"] if b["relleno"])
    tam_etq = int(_EL["pie_foto_etiqueta"]["texto"]["tam_pt"])
    x0 = b0["posicion"]["x"]
    y0 = b0["posicion"]["y"]
    if titulo:
        y0 = _titulo_y_subtitulo(s, X_TEXTO, _EL["titulo_contenido"]["posicion"]["emu"]["y"],
                                 ANCHO_HOJA - 2 * X_TEXTO, titulo, None, "morado") + SEPARACION
    area = Area(x0, y0, ANCHO_HOJA - x0 - X_TEXTO, LIMITE_INF - y0)
    n_col = 1 if len(bloques) <= 4 else 2
    if n_col == 1:
        cols = area.columnas(2, separacion=int(0.4 * EMU_IN))
        grupos = [bloques]
    else:
        cols = area.columnas(2, separacion=int(0.4 * EMU_IN))
        mitad = (len(bloques) + 1) // 2
        grupos = [bloques[:mitad], bloques[mitad:]]
    alto_etq = b0["posicion"]["alto"]
    hueco_etq_texto = int(0.07 * EMU_IN)
    for ci, grupo in enumerate(grupos):
        col = cols[ci]
        ancho_etq = min(b0["posicion"]["ancho"], col.ancho)
        altos = [medir_alto_texto(t, 14, col.ancho) if t else 0 for _, t in grupo]
        total = sum(alto_etq + hueco_etq_texto + h for h in altos)
        libre = col.alto - total
        if libre < 0:
            raise ContenidoDemasiadoLargo("Los bloques no caben: acorta los textos o usa menos bloques.")
        sep = min(int(0.45 * EMU_IN), libre // max(1, len(grupo) - 1)) if len(grupo) > 1 else 0
        y = col.y
        for k, ((etq_txt, cuerpo), h) in enumerate(zip(grupo, altos)):
            color = COLORES_BLOQUES[(ci * len(grupos[0]) + k) % 4]
            if _contar_lineas(etq_txt, tam_etq, ancho_etq - int(0.2 * EMU_IN), negrita=True) > 1:
                raise ContenidoDemasiadoLargo(f"Etiqueta de bloque demasiado larga: {etq_txt!r} (máx. ~28 caracteres).")
            e = _forma(s, "redondeado", col.x, y, ancho_etq, alto_etq, "uceva:bloque_etiqueta", relleno=color)
            e.text_frame.margin_top = e.text_frame.margin_bottom = Emu(0)
            e.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            c_txt = color_texto_sobre(color)
            _escribir(e.text_frame, [etq_txt], tam_etq, c_txt, negrita=True, alineacion="centro", sombra=c_txt == "blanco")
            y += alto_etq + hueco_etq_texto
            if cuerpo:
                tb = _caja_texto(s, col.x, y, col.ancho, h, "uceva:bloque_texto")
                _escribir(tb.text_frame, [cuerpo], 14, "gris_texto")
            y += h + sep
    if imagen_derecha and n_col == 1:
        imagen(s, imagen_derecha, cols[1], borde=True)
    _poner_notas(s, notas)
    return s


def cierre(prs, texto_principal: str = "¡Gracias!", texto_secundario: str | None = None,
           notas: str | None = None):
    """Muestra 11: foto aérea con la banda morada (logo y web). El texto va sobre una capa
    morada semitransparente para que el blanco se lea sobre la foto. Queda al final al guardar."""
    s = _nueva(prs)
    _fondo_imagen(s, "fondo_cierre")
    banda = _ZONAS["banda_cierre"]
    ancho = int(7.0 * EMU_IN)
    ancho_txt = ancho - 2 * int(0.3 * EMU_IN)
    tam_p = 44
    while tam_p > 28 and _contar_lineas(texto_principal, tam_p, ancho_txt, negrita=True) > 1:
        tam_p -= 4
    h_p = medir_alto_texto(texto_principal, tam_p, ancho_txt, negrita=True)
    h_s = medir_alto_texto(texto_secundario, 20, ancho_txt) if texto_secundario else 0
    pad = int(0.25 * EMU_IN)
    alto = h_p + (h_s + int(0.1 * EMU_IN) if texto_secundario else 0) + 2 * pad
    if alto > banda["y"] - 2 * SEPARACION:
        raise ContenidoDemasiadoLargo("Texto de cierre demasiado largo: deja solo una frase corta.")
    y = (banda["y"] - alto) // 2
    x = (ANCHO_HOJA - ancho) // 2
    capa = _forma(s, "redondeado", x, y, ancho, alto, "uceva:capa_cierre", relleno="morado", opacidad=0.85)
    tf = capa.text_frame
    tf.margin_left = tf.margin_right = Emu(int(0.3 * EMU_IN))
    tf.margin_top = tf.margin_bottom = Emu(pad)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    parrafos = [texto_principal] + ([texto_secundario] if texto_secundario else [])
    _escribir(tf, parrafos[:1], tam_p, "blanco", negrita=True, alineacion="centro", sombra=True)
    if texto_secundario:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(6)
        r = p.add_run()
        r.text = texto_secundario
        r.font.size = Pt(20)
        r.font.color.rgb = _rgb("blanco")
        _poner_fuente(r)
        _fin_parrafo(p, 20)
    _estado(prs)["cierres"].append(s.slide_id)
    _poner_notas(s, notas)
    return s


# ---------------------------------------------------------------------------
# Duplicar una muestra tal cual (para quien prefiera editar textos)
# ---------------------------------------------------------------------------

def copiar_muestra(prs, numero: int, notas: str | None = None):
    """Copia exacta de la diapositiva de muestra `numero` (1-11), con fondo e imágenes.
    Después cambia los textos con `reemplazar_texto`. No quedan placeholders vacíos."""
    est = _estado(prs)
    origen = prs.slides.get(est["muestras"][numero - 1])
    s = prs.slides.add_slide(origen.slide_layout)
    _quitar_marcadores(s)
    arbol = s.shapes._spTree
    for el in origen.shapes._spTree.iterchildren():
        if etree.QName(el).localname in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp"):
            arbol.append(copy.deepcopy(el))
    bg = origen._element.find(qn("p:cSld")).find(qn("p:bg"))
    if bg is not None:
        s._element.find(qn("p:cSld")).insert(0, copy.deepcopy(bg))
    # Re-crear las relaciones (imágenes) con los mismos rId de la muestra.
    nuevos = {}
    for el in s._element.iter():
        for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
            rid = el.get(attr)
            if rid and rid in origen.part.rels:
                if rid not in nuevos:
                    rel = origen.part.rels[rid]
                    nuevos[rid] = s.part.relate_to(rel.target_part, rel.reltype)
                el.set(attr, nuevos[rid])
    if numero == len(est["muestras"]):
        est["cierres"].append(s.slide_id)
    _poner_notas(s, notas)
    return s


def reemplazar_texto(slide, viejo: str, nuevo: str) -> int:
    """Cambia el texto de los runs que contienen `viejo` (conserva el formato). Devuelve cuántos cambió."""
    n = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if viejo in r.text:
                    r.text = r.text.replace(viejo, nuevo)
                    n += 1
    return n


# ---------------------------------------------------------------------------
# Elementos de contenido
# ---------------------------------------------------------------------------

def texto(slide, area: Area, contenido: str | list, tam: float = 16, color: str = "gris_texto",
          negrita: bool = False, alineacion: str = "izquierda", vinetas: bool = False,
          tam_minimo: float | None = None, espacio_parrafos_pt: float = 6) -> Area:
    """Texto dentro de `area`, con el alto justo. Si no cabe, baja la letra hasta el mínimo
    (14 pt) y si aun así no cabe lanza ContenidoDemasiadoLargo. `**así**` va en negrita.
    Con vinetas=True, cada elemento de la lista es una viñeta; (texto, 1) es una subviñeta."""
    parrafos = contenido if isinstance(contenido, list) else [contenido]
    if color in ("amarillo", "cian"):
        raise ValueError("No uses amarillo ni cian para texto sobre blanco (no se lee).")
    tam_minimo = tam_minimo or _TAM["cuerpo_minimo"]
    sangria = int(0.25 * EMU_IN) if vinetas else 0
    t = tam
    while True:
        planos = [p[0] if isinstance(p, tuple) else p for p in parrafos]
        alto = medir_alto_texto(planos, t, area.ancho, negrita, espacio_parrafos_pt=espacio_parrafos_pt,
                                sangria_emu=sangria * (2 if any(isinstance(p, tuple) and p[1] for p in parrafos) else 1))
        if alto <= area.alto or t <= tam_minimo:
            break
        t -= 1
    if alto > area.alto:
        raise ContenidoDemasiadoLargo(f"El texto no cabe ni a {t} pt ({alto / EMU_IN:.2f}\" de {area.alto / EMU_IN:.2f}\"). "
                                      "Divídelo en dos diapositivas o resúmelo.")
    tb = _caja_texto(slide, area.x, area.y, area.ancho, alto, "uceva:texto")
    _escribir(tb.text_frame, parrafos, t, color, negrita=negrita, alineacion=alineacion, vinetas=vinetas,
              espacio_parrafos_pt=espacio_parrafos_pt)
    return area.con_alto(alto)


def _alto_tarjeta(titulo, cuerpo, ancho, tam_cuerpo):
    pad = int(0.15 * EMU_IN)
    alto_etq = int(0.42 * EMU_IN)
    cuerpo_l = cuerpo if isinstance(cuerpo, list) else ([cuerpo] if cuerpo else [])
    h = medir_alto_texto(cuerpo_l, tam_cuerpo, ancho - 2 * pad, espacio_parrafos_pt=4) if cuerpo_l else 0
    return pad + alto_etq + (int(0.1 * EMU_IN) + h if cuerpo_l else 0) + pad, h


def tarjeta(slide, x: int, y: int, ancho: int, titulo: str, cuerpo: str | list | None = None,
            color: str = "morado", alto: int | None = None, tam_cuerpo: float = 16) -> Area:
    """Tarjeta blanca con borde gris claro, cabecera de color con `titulo` y texto debajo.
    El alto se ajusta al contenido (sin hueco vacío abajo) salvo que pases `alto`."""
    pad = int(0.15 * EMU_IN)
    alto_etq = int(0.42 * EMU_IN)
    necesario, h_cuerpo = _alto_tarjeta(titulo, cuerpo, ancho, tam_cuerpo)
    alto = max(alto or 0, necesario)
    fondo = _forma(slide, "redondeado", x, y, ancho, alto, "uceva:tarjeta", relleno="blanco", linea="gris_claro", ancho_linea_pt=1)
    fondo.adjustments[0] = 0.06
    tam_t = 16
    while tam_t > 14 and _contar_lineas(titulo, tam_t, ancho - 2 * pad - int(0.2 * EMU_IN), negrita=True) > 1:
        tam_t -= 1
    if _contar_lineas(titulo, tam_t, ancho - 2 * pad - int(0.2 * EMU_IN), negrita=True) > 1:
        raise ContenidoDemasiadoLargo(f"Título de tarjeta demasiado largo para su ancho: {titulo!r}")
    e = _forma(slide, "redondeado", x + pad, y + pad, ancho - 2 * pad, alto_etq, "uceva:tarjeta_titulo", relleno=color)
    e.text_frame.margin_top = e.text_frame.margin_bottom = Emu(0)
    e.text_frame.margin_left = e.text_frame.margin_right = Emu(int(0.1 * EMU_IN))
    e.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    c_txt = color_texto_sobre(color)
    _escribir(e.text_frame, [titulo], tam_t, c_txt, negrita=True, alineacion="centro", sombra=c_txt == "blanco")
    if cuerpo:
        cuerpo_l = cuerpo if isinstance(cuerpo, list) else [cuerpo]
        tb = _caja_texto(slide, x + pad, y + pad + alto_etq + int(0.1 * EMU_IN), ancho - 2 * pad, h_cuerpo, "uceva:tarjeta_texto")
        _escribir(tb.text_frame, cuerpo_l, tam_cuerpo, "gris_texto", espacio_parrafos_pt=4)
    return Area(x, y, ancho, alto)


def fila_de_tarjetas(slide, area: Area, tarjetas: list[tuple[str, str]], colores: list[str] | None = None,
                     tam_cuerpo: float = 16) -> Area:
    """Tarjetas del mismo alto en una fila (2 a 4 recomendado). Colores por defecto: los de la marca en orden."""
    n = len(tarjetas)
    if not 1 <= n <= 5:
        raise ContenidoDemasiadoLargo("Usa de 1 a 5 tarjetas por fila (mejor 2-4).")
    cols = area.columnas(n)
    colores = colores or [COLORES_BLOQUES[i % 4] for i in range(n)]
    alto = max(_alto_tarjeta(t, c, cols[i].ancho, tam_cuerpo)[0] for i, (t, c) in enumerate(tarjetas))
    if alto > area.alto:
        raise ContenidoDemasiadoLargo("Las tarjetas no caben: acorta los textos o usa menos tarjetas.")
    for col, (t, c), color in zip(cols, tarjetas, colores):
        tarjeta(slide, col.x, col.y, col.ancho, t, c, color=color, alto=alto, tam_cuerpo=tam_cuerpo)
    return area.con_alto(alto)


def cuadricula(slide, area: Area, tarjetas: list[tuple[str, str]], columnas: int = 2,
               colores: list[str] | None = None, tam_cuerpo: float = 16) -> Area:
    """Tarjetas en cuadrícula de `columnas`; cada fila con el alto de su tarjeta más alta."""
    colores = colores or [COLORES_BLOQUES[i % 4] for i in range(len(tarjetas))]
    y = area.y
    for i in range(0, len(tarjetas), columnas):
        fila = tarjetas[i:i + columnas]
        sub = Area(area.x, y, area.ancho, area.abajo - y)
        cols = sub.columnas(columnas)
        alto = max(_alto_tarjeta(t, c, cols[k].ancho, tam_cuerpo)[0] for k, (t, c) in enumerate(fila))
        if y + alto > area.abajo:
            raise ContenidoDemasiadoLargo("La cuadrícula no cabe: divide las tarjetas en dos diapositivas.")
        for k, (t, c) in enumerate(fila):
            tarjeta(slide, cols[k].x, y, cols[k].ancho, t, c, color=colores[i + k], alto=alto, tam_cuerpo=tam_cuerpo)
        y += alto + SEPARACION
    return area.con_alto(y - SEPARACION - area.y)


def dividir_tabla(filas: list[list], max_filas: int) -> list[list[list]]:
    """Parte una tabla larga en trozos de como mucho `max_filas` filas de datos, repitiendo el
    encabezado. Reparte parejo (20 filas con máx. 8 → 7 + 7 + 6, no 8 + 8 + 4)."""
    enc, datos = filas[0], filas[1:]
    n = max(1, -(-len(datos) // max_filas))
    tam, resto = divmod(len(datos), n)
    trozos, i = [], 0
    for k in range(n):
        j = i + tam + (1 if k < resto else 0)
        trozos.append([enc] + datos[i:j])
        i = j
    return trozos


def tabla(slide, area: Area, filas: list[list], anchos: list[float] | None = None, tam: float = 16,
          alineaciones: list[str] | None = None) -> Area:
    """Tabla con encabezado morado (texto blanco) y filas alternas morado suave / blanco.
    `filas[0]` es el encabezado. `anchos`: proporciones de columna. Si no cabe, lanza
    ContenidoDemasiadoLargo (usa `dividir_tabla`)."""
    n_f, n_c = len(filas), len(filas[0])
    anchos = anchos or [1] * n_c
    total = sum(anchos)
    anchos_emu = [int(area.ancho * a / total) for a in anchos]
    mx, my = int(0.08 * EMU_IN), int(0.05 * EMU_IN)
    alturas = []
    for i, fila in enumerate(filas):
        h = max(medir_alto_texto(str(c), tam, w - 2 * mx, negrita=(i == 0)) for c, w in zip(fila, anchos_emu))
        alturas.append(h + 2 * my)
    if sum(alturas) > area.alto:
        caben, acumulado = 0, alturas[0]
        for h in alturas[1:]:
            if acumulado + h > area.alto:
                break
            acumulado += h
            caben += 1
        raise ContenidoDemasiadoLargo(f"La tabla ({n_f - 1} filas de datos) no cabe en {area.alto / EMU_IN:.2f}\"; "
                                      f"caben unas {caben}. Usa dividir_tabla(filas, {max(1, caben)}) y una "
                                      "diapositiva por trozo.")
    gf = slide.shapes.add_table(n_f, n_c, Emu(area.x), Emu(area.y), Emu(area.ancho), Emu(sum(alturas)))
    gf.name = "uceva:tabla"
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    estilo = tblPr.find(qn("a:tableStyleId"))
    if estilo is None:
        estilo = etree.SubElement(tblPr, qn("a:tableStyleId"))
    estilo.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"    # "Sin estilo, sin cuadrícula": todo lo pone la librería
    for j, w in enumerate(anchos_emu):
        tbl.columns[j].width = Emu(w)
    alineaciones = alineaciones or ["izquierda"] * n_c
    for i, fila in enumerate(filas):
        tbl.rows[i].height = Emu(alturas[i])
        for j, valor in enumerate(fila):
            celda = tbl.cell(i, j)
            celda.margin_left = celda.margin_right = Emu(mx)
            celda.margin_top = celda.margin_bottom = Emu(my)
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
            fondo = "morado" if i == 0 else ("morado_suave" if i % 2 == 0 else "blanco")
            celda.fill.solid()
            celda.fill.fore_color.rgb = _rgb(fondo)
            _escribir(celda.text_frame, [str(valor)], tam, "blanco" if i == 0 else "gris_texto",
                      negrita=(i == 0), alineacion=alineaciones[j])
            tcPr = celda._tc.get_or_add_tcPr()
            ln = etree.Element(qn("a:lnB"), w="9525")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", COLORES["gris_claro"])
            tcPr.insert(0, ln)
    return area.con_alto(sum(alturas))


def bloque_codigo(slide, area: Area, codigo: str, tam: float = 14, lenguaje: str | None = None) -> Area:
    """Código en Courier New, texto blanco sobre morado muy oscuro. Las líneas no se parten:
    si alguna es muy larga baja la letra (mín. 12 pt) y si no cabe lanza ContenidoDemasiadoLargo."""
    lineas = codigo.expandtabs(4).rstrip("\n").split("\n")
    pad_x, pad_y = int(0.2 * EMU_IN), int(0.15 * EMU_IN)
    extra_etq = int(0.3 * EMU_IN) if lenguaje else 0
    t = tam
    while True:
        ancho_max = max(_ancho_pt(l, t, FUENTE_CODIGO) for l in lineas) * EMU_PT / 0.97
        alto = int(len(lineas) * t * INTERLINEADO * EMU_PT) + 2 * pad_y + extra_etq
        if (ancho_max <= area.ancho - 2 * pad_x and alto <= area.alto) or t <= _TAM["codigo_minimo"]:
            break
        t -= 1
    if ancho_max > area.ancho - 2 * pad_x or alto > area.alto:
        raise ContenidoDemasiadoLargo(f"El código ({len(lineas)} líneas) no cabe. Muestra solo el fragmento clave "
                                      "(máx. ~14 líneas de ~70 caracteres) o pártelo en dos diapositivas.")
    fondo = _forma(slide, "redondeado", area.x, area.y, area.ancho, alto, "uceva:codigo", relleno="morado_oscuro")
    fondo.adjustments[0] = 0.04
    tf = fondo.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Emu(pad_x)
    tf.margin_top = Emu(pad_y + extra_etq)
    tf.margin_bottom = Emu(pad_y)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.clear()
    for i, l in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = l if l else " "
        r.font.size = Pt(t)
        r.font.color.rgb = _rgb("blanco")
        _poner_fuente(r, FUENTE_CODIGO)
        _fin_parrafo(p, t, FUENTE_CODIGO)
    if lenguaje:
        tb = _caja_texto(slide, area.x + pad_x, area.y + int(0.08 * EMU_IN), area.ancho - 2 * pad_x, int(0.25 * EMU_IN), "uceva:codigo_lenguaje")
        _escribir(tb.text_frame, [lenguaje], 12, "cian", negrita=True)
    return area.con_alto(alto)


def imagen(slide, ruta: str | Path, area: Area, alinear: str = "centro", borde: bool = False) -> Area:
    """Imagen dentro de `area` sin deformarla (la encaja y la centra). borde=True: borde blanco y sombra."""
    with Image.open(ruta) as im:
        prop = im.width / im.height
    ancho, alto = area.ancho, int(area.ancho / prop)
    if alto > area.alto:
        alto, ancho = area.alto, int(area.alto * prop)
    x = {"izquierda": area.x, "centro": area.x + (area.ancho - ancho) // 2, "derecha": area.derecha - ancho}[alinear]
    y = area.y + (area.alto - alto) // 2
    pic = slide.shapes.add_picture(str(ruta), Emu(x), Emu(y), Emu(ancho), Emu(alto))
    pic.name = "uceva:imagen"
    if borde:
        _borde_y_sombra(pic)
    return Area(x, y, ancho, alto)


def numero_grande(slide, x: int, y: int, ancho: int, numero: str, descripcion: str,
                  color: str = "morado", tam: float = 54) -> Area:
    """Cifra grande (54 pt negrita) con una descripción corta debajo (14 pt), centradas."""
    if color in ("amarillo", "cian"):
        raise ValueError("Para cifras usa morado o rosado (amarillo y cian no se leen sobre blanco).")
    alto_n = int(tam * INTERLINEADO * EMU_PT)
    tb = _caja_texto(slide, x, y, ancho, alto_n, "uceva:numero")
    _escribir(tb.text_frame, [numero], tam, color, negrita=True, alineacion="centro")
    alto_d = medir_alto_texto(descripcion, 14, ancho)
    td = _caja_texto(slide, x, y + alto_n, ancho, alto_d, "uceva:numero_texto")
    _escribir(td.text_frame, [descripcion], 14, "gris_texto", alineacion="centro")
    return Area(x, y, ancho, alto_n + alto_d)


def fila_de_numeros(slide, area: Area, cifras: list[tuple[str, str]], colores: list[str] | None = None) -> Area:
    """2 a 4 cifras grandes en fila, en el tercio superior de `area`."""
    colores = colores or ["morado", "rosado"] * 3
    cols = area.columnas(len(cifras))
    alto = max(int(54 * INTERLINEADO * EMU_PT) + medir_alto_texto(d, 14, c.ancho) for c, (_, d) in zip(cols, cifras))
    if alto > area.alto:
        raise ContenidoDemasiadoLargo("Las cifras no caben: acorta las descripciones.")
    y = area.y + (area.alto - alto) // 3
    for c, (n, d), color in zip(cols, cifras, colores):
        numero_grande(slide, c.x, y, c.ancho, n, d, color=color)
    return Area(area.x, y, area.ancho, alto)
