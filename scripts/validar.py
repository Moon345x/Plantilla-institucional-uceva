#!/usr/bin/env python3
"""Revisa un .pptx hecho con la plantilla UCEVA.

    python scripts/validar.py salida/ejemplo.pptx [--estricto]

Revisa:
  - que abre, que no hay partes huérfanas ni relaciones rotas;
  - diapositivas de muestra que quedaron sin usar y textos de ejemplo de la plantilla;
  - marcadores (placeholders) vacíos;
  - texto que desborda su forma o la hoja (estimado con el tamaño de letra y la caja);
  - formas fuera de la hoja;
  - tamaños de letra por debajo de los mínimos;
  - formas dentro de las zonas protegidas (logo, etiqueta, barra lateral, banda del cierre);
  - colores fuera de spec/tokens.json y fuentes distintas de Arial / Arial Black / Courier New.

Termina con un resumen (OK / avisos / errores). Sale con código 1 si hay errores
(con --estricto, también si hay avisos).
"""
from __future__ import annotations

import argparse
import hashlib
import posixpath
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "lib"))
import uceva_pptx as u  # noqa: E402

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
R_EMBED = f"{{{NS['r']}}}embed"

T = u.TOKENS
W, H = u.ANCHO_HOJA, u.ALTO_HOJA
ZONAS = T["zonas_protegidas"]
SHA = {v["sha1"]: k for k, v in T["imagenes"].items()}
PALETA = set(u.COLORES.values())
FUENTES_OK = {T["fuentes"]["texto"], T["fuentes"]["titulo_portada"], T["fuentes"]["codigo"]}
MIN_CUERPO = T["tamanos_pt"]["cuerpo_minimo"]
MIN_PIE = T["tamanos_pt"]["pie_minimo"]
TOL = int(0.02 * 914400)
# Formas que pueden usar 12-13 pt (etiquetas y pies) o salirse de la hoja (sangrado del marco).
NOMBRES_PIE = ("uceva:codigo_lenguaje", "uceva:pie", "uceva:etiqueta", "uceva:autor")
NOMBRES_SANGRADO = ("uceva:marco_portada",)
TEXTOS_PROMPT = ("Haga clic para", "Click to add", "Haz clic para")


class Informe:
    def __init__(self):
        self.errores, self.avisos = [], []

    def error(self, donde, msg):
        self.errores.append(f"[{donde}] {msg}")

    def aviso(self, donde, msg):
        self.avisos.append(f"[{donde}] {msg}")


# --- estructura del paquete ---------------------------------------------------

def revisar_paquete(ruta: Path, inf: Informe):
    z = zipfile.ZipFile(ruta)
    nombres = set(z.namelist())
    alcanzables, rotas = set(), []

    def rels_de(parte):
        d, f = posixpath.split(parte)
        return posixpath.join(d, "_rels", f + ".rels")

    pendientes = [""]
    vistos = set()
    while pendientes:
        parte = pendientes.pop()
        if parte in vistos:
            continue
        vistos.add(parte)
        archivo_rels = "_rels/.rels" if parte == "" else rels_de(parte)
        if archivo_rels not in nombres:
            continue
        alcanzables.add(archivo_rels)
        base = posixpath.dirname(parte)
        for r in etree.fromstring(z.read(archivo_rels)).findall("rel:Relationship", NS):
            if r.get("TargetMode") == "External":
                continue
            destino = r.get("Target")
            destino = destino.lstrip("/") if destino.startswith("/") else posixpath.normpath(posixpath.join(base, destino))
            if destino not in nombres:
                rotas.append(f"{parte or '(paquete)'} → {destino}")
                continue
            alcanzables.add(destino)
            pendientes.append(destino)
    # Diapositivas quitadas de la lista sin soltar su relación (falta drop_rel).
    pres_rels = {r.get("Id"): r.get("Target") for r in etree.fromstring(z.read("ppt/_rels/presentation.xml.rels")).findall("rel:Relationship", NS)
                 if r.get("Type").endswith("/slide")}
    en_lista = {e.get(f"{{{NS['r']}}}id") for e in etree.fromstring(z.read("ppt/presentation.xml")).iterfind(".//p:sldId", NS)}
    for rid in sorted(set(pres_rels) - en_lista):
        inf.error("paquete", f"{pres_rels[rid]} no está en la lista de diapositivas pero sigue en el archivo "
                             "(se borró sin drop_rel): parte huérfana")
    huerfanas = sorted(n for n in nombres - alcanzables if n != "[Content_Types].xml" and not n.endswith("/"))
    for h in huerfanas:
        inf.error("paquete", f"parte huérfana (nadie la usa): {h}")
    for r in rotas:
        inf.error("paquete", f"relación rota: {r}")


# --- firmas de las muestras de la plantilla ----------------------------------

def firma(slide) -> tuple:
    bg = slide._element.find("p:cSld/p:bg/p:bgPr/a:blipFill/a:blip", NS)
    sha_bg = hashlib.sha1(slide.part.related_part(bg.get(R_EMBED)).blob).hexdigest() if bg is not None else None
    formas = tuple(sorted((sh.left, sh.top, sh.width, sh.height, sh.text_frame.text if sh.has_text_frame else "",
                           tuple(c.get("val") for c in sh._element.iter(f"{{{NS['a']}}}srgbClr")))
                          for sh in slide.shapes))
    return sha_bg, formas


def firmas_muestras() -> dict:
    from pptx import Presentation
    prs = Presentation(str(u.PLANTILLA_POR_DEFECTO))
    return {firma(s): i for i, s in enumerate(prs.slides, start=1)}


# --- texto --------------------------------------------------------------------

def tam_run(r, defecto=18.0):
    sz = r.get("sz")
    return int(sz) / 100 if sz else defecto


def parrafos_de(txBody):
    """[(texto, tam_max, negrita, fuente)] por párrafo."""
    res = []
    for p in txBody.findall("a:p", NS):
        runs = p.findall("a:r", NS)
        t = "".join(r.findtext("a:t", default="", namespaces=NS) for r in runs)
        if runs:
            rprs = [r.find("a:rPr", NS) for r in runs]
            tams = [tam_run(x) if x is not None else 18.0 for x in rprs]
            neg = any(x is not None and x.get("b") == "1" for x in rprs)
            lat = rprs[0].find("a:latin", NS) if rprs[0] is not None else None
            fuente = lat.get("typeface") if lat is not None else "Arial"
            res.append((t, max(tams), neg, fuente))
        else:
            end = p.find("a:endParaRPr", NS)
            res.append(("", tam_run(end) if end is not None else 18.0, False, "Arial"))
    return res


def alto_necesario(parrafos, ancho, envolver=True):
    total = 0
    for t, tam, neg, fuente in parrafos:
        fam = fuente if fuente in ("Arial", "Arial Black", "Courier New") else "Arial"
        lineas = u._contar_lineas(t, tam, ancho, fam, neg) if (t and envolver) else 1
        total += lineas * tam * u.INTERLINEADO * u.EMU_PT
    return int(total)


def margenes(bodyPr):
    g = lambda k, d: int(bodyPr.get(k, d)) if bodyPr is not None else d  # noqa: E731
    return g("lIns", 91440), g("tIns", 45720), g("rIns", 91440), g("bIns", 45720)


def interseca(a, b) -> bool:
    return a[0] < b[0] + b[2] - TOL and b[0] < a[0] + a[2] - TOL and a[1] < b[1] + b[3] - TOL and b[1] < a[1] + a[3] - TOL


def contiene(a, b) -> bool:
    return a[0] <= b[0] + TOL and a[1] <= b[1] + TOL and b[0] + b[2] <= a[0] + a[2] + TOL and b[1] + b[3] <= a[1] + a[3] + TOL


def caja(z):
    return (z["x"], z["y"], z["ancho"], z["alto"])


# --- revisión por diapositiva ----------------------------------------------------

def revisar_diapositiva(n, s, inf: Informe, muestras: dict):
    donde = f"diapositiva {n}"
    el = s._element
    # fondo
    bg = el.find("p:cSld/p:bg/p:bgPr/a:blipFill/a:blip", NS)
    fondo = SHA.get(hashlib.sha1(s.part.related_part(bg.get(R_EMBED)).blob).hexdigest()) if bg is not None else "fondo_contenido"
    if bg is None and s.slide_layout.name not in ("Diapositiva de título", "En blanco"):
        inf.aviso(donde, f"usa el diseño '{s.slide_layout.name}' (la plantilla usa 'Diapositiva de título')")

    f = firma(s)
    # Las muestras que son solo una imagen de fondo (bienvenida, cierre) se usan tal cual a propósito.
    if f in muestras and not (f[0] and not f[1]):
        inf.aviso(donde, f"es idéntica a la muestra {muestras[f]} de la plantilla. ¿Quedó sin usar?")

    # zonas protegidas según el tipo de fondo
    zonas = []
    if fondo in ("fondo_contenido", None):
        zonas.append(("logo", caja(ZONAS["logo_contenido"]), ()))
    elif fondo in ("fondo_portada", "fondo_bienvenidos"):
        zonas.append(("logo", caja(ZONAS["logo_portada"]), ()))
        if fondo == "fondo_portada":
            zonas.append(("arte 'Bienvenidos'", caja(ZONAS["arte_bienvenidos_portada"]), ()))
    elif fondo == "fondo_barra_lateral":
        zonas.append(("barra lateral", caja(ZONAS["barra_lateral"]), ("uceva:barra", "uceva:logo")))
    elif fondo == "fondo_cierre":
        zonas.append(("banda del cierre", caja(ZONAS["banda_cierre"]), ()))
    etiquetas = [sh for sh in s.shapes if sh.name == "uceva:etiqueta"]
    for e in etiquetas:
        zonas.append(("etiqueta de facultad", (e.left, e.top, e.width, e.height), ("uceva:etiqueta",)))

    rellenas = []   # formas con relleno ya dibujadas (el texto encima no está "sobre blanco")
    textos = []     # (nombre, caja real con el alto que de verdad ocupa el texto)
    for sh in s.shapes:
        nombre = sh.name
        d = f"{donde} · {nombre}"
        if sh.is_placeholder:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                inf.error(d, "marcador (placeholder) vacío: bórralo")
                continue
        if sh.left is None:
            continue
        rect = (sh.left, sh.top, sh.width, sh.height)
        # fuera de la hoja
        if nombre not in NOMBRES_SANGRADO and (sh.left < -TOL or sh.top < -TOL or sh.left + sh.width > W + TOL or sh.top + sh.height > H + TOL):
            inf.error(d, "se sale de la hoja")
        # zonas protegidas
        for zona, z, duenos in zonas:
            if nombre in duenos or any(nombre == dd for dd in duenos):
                continue
            if fondo == "fondo_barra_lateral" and nombre in ("uceva:barra", "uceva:logo"):
                continue
            if interseca(rect, z):
                inf.error(d, f"invade la zona protegida: {zona}")
        # texto
        bodies = []
        if sh.has_text_frame:
            bodies.append((sh._element.find(".//p:txBody", NS), sh.width, sh.height, sh.top))
        if getattr(sh, "has_table", False) and sh.has_table:
            revisar_tabla(d, sh, inf)
        for txBody, ancho, alto, top in bodies:
            if txBody is None:
                continue
            encima = any(r[0] <= sh.left and r[1] <= sh.top and sh.left + sh.width <= r[0] + r[2]
                         and sh.top + sh.height <= r[1] + r[3] for r in rellenas)
            alto_real = revisar_texto(d, nombre, sh, txBody, ancho, alto, top, inf, sobre_color=_tiene_relleno(sh) or encima)
            if alto_real:
                textos.append((nombre, (sh.left, sh.top, sh.width, alto_real), _tiene_relleno(sh)))
        if _tiene_relleno(sh):
            rellenas.append(rect)

    # Textos que se pisan (p. ej. un título largo que baja hasta el subtítulo). Solo se permite
    # texto dentro de una forma con relleno (etiqueta del código, texto sobre una caja de color).
    for i, (n1, a, r1) in enumerate(textos):
        for n2, b, r2 in textos[i + 1:]:
            if interseca(a, b) and not ((r1 and contiene(a, b)) or (r2 and contiene(b, a))):
                inf.error(f"{donde} · {n1}", f"el texto se pisa con «{n2}»")


def _tiene_relleno(sh) -> bool:
    spPr = sh._element.find("p:spPr", NS)
    return spPr is not None and spPr.find("a:solidFill", NS) is not None


def revisar_texto(d, nombre, sh, txBody, ancho, alto, top, inf, sobre_color=False) -> int:
    """Revisa un cuadro de texto. Devuelve el alto que ocupa de verdad (0 si está vacío)."""
    texto_total = "".join(t.text or "" for t in txBody.iter(f"{{{NS['a']}}}t"))
    if not texto_total.strip():
        return 0
    for ej in u.TEXTOS_DE_EJEMPLO:
        if ej.lower() in texto_total.lower():
            inf.error(d, f"quedó texto de ejemplo de la plantilla: «{ej}»")
    for pr in TEXTOS_PROMPT:
        if pr in texto_total:
            inf.error(d, f"quedó texto de marcador: «{pr}…»")
    # tamaños, fuentes y colores
    for r in txBody.iter(f"{{{NS['a']}}}r"):
        rpr = r.find("a:rPr", NS)
        t = r.findtext("a:t", default="", namespaces=NS)
        if not t.strip():
            continue
        tam = tam_run(rpr) if rpr is not None else 18.0
        if tam < MIN_PIE:
            inf.error(d, f"letra de {tam:g} pt (mínimo {MIN_PIE:g}): «{t[:30]}»")
        elif tam < MIN_CUERPO and not nombre.startswith(NOMBRES_PIE):
            inf.aviso(d, f"letra de {tam:g} pt; el cuerpo debe tener al menos {MIN_CUERPO:g}: «{t[:30]}»")
        lat = rpr.find("a:latin", NS) if rpr is not None else None
        fuente = lat.get("typeface") if lat is not None else None
        if fuente is None:
            if not sh.is_placeholder:
                inf.aviso(d, f"texto sin fuente declarada (heredará {T['fuentes']['fuente_del_tema']}): «{t[:30]}»")
        elif fuente not in FUENTES_OK:
            inf.aviso(d, f"fuente «{fuente}» fuera de la marca (usa Arial)")
        elif fuente == T["fuentes"]["titulo_portada"] and nombre != "uceva:titulo_portada":
            inf.aviso(d, "Arial Black solo se usa en el título de la portada")
        clr = rpr.find("a:solidFill/a:srgbClr", NS) if rpr is not None else None
        if clr is not None and not sobre_color and clr.get("val").upper() in (u.COLORES["amarillo"], u.COLORES["cian"]):
            inf.aviso(d, "texto amarillo o cian sobre fondo claro: casi no se lee")
    # desborde
    bodyPr = txBody.find("a:bodyPr", NS)
    l, t_, r_, b = margenes(bodyPr)
    envolver = bodyPr is None or bodyPr.get("wrap") != "none"
    pars = parrafos_de(txBody)
    necesario = alto_necesario(pars, ancho - l - r_, envolver) + t_ + b
    if not envolver:
        ancho_max = max((u._ancho_pt(p[0], p[1], p[3] if p[3] in ("Arial", "Courier New") else "Arial", p[2]) for p in pars), default=0)
        if ancho_max * u.EMU_PT > ancho - l - r_:
            inf.error(d, "una línea es más ancha que la caja (texto sin ajuste de línea)")
    crece = bodyPr is not None and bodyPr.find("a:spAutoFit", NS) is not None
    if crece:
        if top + necesario > H + TOL:
            inf.error(d, "el texto hace crecer la caja más allá del borde inferior de la hoja")
        elif necesario > alto + int(0.1 * 914400):
            inf.aviso(d, f"la caja crecerá a ~{necesario / 914400:.2f}\" al editarla (ahora mide {alto / 914400:.2f}\")")
    elif necesario > alto + int(0.1 * 914400):
        inf.error(d, f"el texto desborda la forma (necesita ~{necesario / 914400:.2f}\", tiene {alto / 914400:.2f}\")")
    elif necesario > alto + TOL:
        inf.aviso(d, f"el texto queda muy justo (necesita ~{necesario / 914400:.2f}\", tiene {alto / 914400:.2f}\")")
    return max(alto, necesario)


def revisar_tabla(d, sh, inf):
    tbl = sh.table
    total = 0
    for i, fila in enumerate(tbl.rows):
        alto_fila = fila.height
        for j, celda in enumerate(fila.cells):
            tx = celda._tc.find("a:txBody", NS)
            if tx is None:
                continue
            ancho = tbl.columns[j].width - (celda.margin_left or 0) - (celda.margin_right or 0)
            nec = alto_necesario(parrafos_de(tx), ancho) + (celda.margin_top or 0) + (celda.margin_bottom or 0)
            alto_fila = max(alto_fila, nec)
            revisar_texto(f"{d} celda {i + 1},{j + 1}", sh.name, sh, tx, tbl.columns[j].width, 10 ** 9, 0, Informe(),
                          sobre_color=True)
            for r in tx.iter(f"{{{NS['a']}}}r"):
                rpr = r.find("a:rPr", NS)
                tam = tam_run(rpr) if rpr is not None else 18.0
                if tam < MIN_PIE:
                    inf.error(d, f"celda {i + 1},{j + 1} con letra de {tam:g} pt")
        total += alto_fila
    if sh.top + total > H - T["margenes"]["inferior_minimo_emu"] + TOL:
        inf.error(d, f"la tabla crece hasta {(sh.top + total) / 914400:.2f}\" y se sale por abajo: divídela")


def revisar_colores(n, s, inf):
    vistos = set()
    for clr in s._element.iter(f"{{{NS['a']}}}srgbClr"):
        if clr.getparent().tag == f"{{{NS['a']}}}outerShdw":
            continue
        v = clr.get("val").upper()
        if v not in PALETA:
            vistos.add(v)
    for v in sorted(vistos):
        inf.aviso(f"diapositiva {n}", f"color {v} fuera de spec/tokens.json")


def validar(ruta: Path, estricto: bool = False) -> int:
    from pptx import Presentation
    inf = Informe()
    try:
        prs = Presentation(str(ruta))
    except Exception as e:  # noqa: BLE001
        print(f"ERROR: no se pudo abrir {ruta}: {e}")
        return 1
    revisar_paquete(ruta, inf)
    if (prs.slide_width, prs.slide_height) != (W, H):
        inf.error("presentación", f"tamaño {prs.slide_width}×{prs.slide_height} EMU; la plantilla es {W}×{H} (16:9)")
    muestras = firmas_muestras()
    coincidencias = 0
    con_notas = 0
    for n, s in enumerate(prs.slides, start=1):
        if firma(s) in muestras:
            coincidencias += 1
        revisar_diapositiva(n, s, inf, muestras)
        revisar_colores(n, s, inf)
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            con_notas += 1
    if len(prs.slides) >= len(muestras) and coincidencias >= len(muestras):
        inf.error("presentación", "siguen todas las diapositivas de muestra: usa guardar() de la librería para quitarlas")

    print(f"Validando {ruta}  ({len(prs.slides)} diapositivas, {con_notas} con notas del orador)\n")
    for e in inf.errores:
        print(f"  ERROR  {e}")
    for a in inf.avisos:
        print(f"  AVISO  {a}")
    estado = "ERRORES" if inf.errores else ("AVISOS" if inf.avisos else "OK")
    print(f"\nResumen: {estado} · {len(inf.errores)} errores · {len(inf.avisos)} avisos")
    return 1 if inf.errores or (estricto and inf.avisos) else 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--estricto", action="store_true", help="los avisos también hacen fallar")
    a = ap.parse_args()
    sys.exit(validar(a.pptx, a.estricto))


if __name__ == "__main__":
    main()
