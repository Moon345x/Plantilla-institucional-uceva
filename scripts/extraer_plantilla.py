#!/usr/bin/env python3
"""Lee la plantilla .pptx UNA vez y deja todo como archivos legibles.

Genera (dentro de la carpeta de salida, por defecto la raíz del repo):
  spec/tokens.json    colores, fuentes, tamaños, márgenes y zonas protegidas
  spec/layouts.json   diseños (layouts) del archivo y para qué sirve cada uno
  spec/slides.json    cada diapositiva de muestra con sus formas y estilos
  assets/img/         imágenes de la plantilla con nombres claros
  assets/previews/    un PNG por diapositiva de muestra (si hay motor de render)

Uso:
  python scripts/extraer_plantilla.py                    # todo en el repo
  python scripts/extraer_plantilla.py --salida /tmp/x    # en otra carpeta
  python scripts/extraer_plantilla.py --sin-previews     # sin renderizar

No modifica el .pptx original: lo abre en modo lectura.
"""
from __future__ import annotations

import argparse
import colorsys
import hashlib
import io
import json
import posixpath
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops

RAIZ = Path(__file__).resolve().parent.parent
PPTX_POR_DEFECTO = RAIZ / "plantilla" / "PLANTILLA-PRESENTACION-INSTITUCIONAL-55-anos_VF.pptx"

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
EMU_POR_PULGADA = 914400
EMU_POR_PT = 12700

# --- Nombres y descripciones (texto editorial; los datos numéricos salen del archivo) ---

NOMBRES_COLOR = {
    "6000AF": "morado",
    "FF0066": "rosado",
    "FFC403": "amarillo",
    "00E8FC": "cian",
    "FFFFFF": "blanco",
    "000000": "negro",
}
USO_COLOR = {
    "morado": "Color institucional principal: títulos, etiquetas, cajas, barra lateral.",
    "rosado": "Acento: etiqueta de facultad, barra lateral, bloques.",
    "amarillo": "Acento: etiquetas y bloques. No usar como color de texto sobre blanco.",
    "cian": "Acento: marco punteado de la portada y bloques. No usar como texto sobre blanco.",
    "blanco": "Texto sobre fondos de color; borde de fotos.",
    "negro": "Solo sombras (con transparencia). No usar para texto.",
    "gris_texto": "Texto corriente (cuerpo, descripciones, pies).",
    "gris_claro": "Color 'Fondo 2' del tema. Bordes finos y separadores suaves.",
}

# Rol de cada imagen según dónde la usa la plantilla -> nombre de archivo claro.
ROL_IMAGEN = {
    "fondo_patron": "fondo_contenido",          # fondo del patrón: logo pequeño + figuras grises
    "fondo_slide1": "fondo_bienvenidos",         # portada "Bienvenidos"
    "fondo_slide2": "fondo_portada",             # portada con título
    "fondo_slide7": "fondo_barra_lateral",       # fondo sin logo (el logo va en la barra)
    "imagen_slide7_Imagen 6": "logo_55_uceva_blanco",
    "imagen_slide9_Imagen 4": "foto_ejemplo_campus",
    "fondo_slide11": "fondo_cierre",
}

DESCRIPCION_SLIDE = {
    1: ("bienvenida", "Portada 'Bienvenidos'. Solo imagen de fondo (logo grande + arte 'Bienvenidos'). Para abrir un evento o una clase.", "bienvenida(prs)"),
    2: ("portada", "Portada con título: caja morada con el título (Arial Black, blanco), marco punteado cian y subtítulo morado. El fondo trae el logo y un 'Bienvenidos' pequeño abajo a la izquierda.", "portada(prs, titulo, subtitulo, autor)"),
    3: ("en_blanco", "En blanco con el logo del patrón arriba a la izquierda. Usa el diseño 'En blanco'. Para contenido libre.", "diapositiva_en_blanco(prs)"),
    4: ("contenido_etiqueta_amarilla", "Contenido con etiqueta de facultad AMARILLA arriba a la derecha + subtítulo amarillo y texto corto. (La librería deja el título en morado porque el amarillo sobre blanco no se lee bien.)", "diapositiva_contenido(prs, etiqueta, titulo, color='amarillo')"),
    5: ("contenido_etiqueta_rosada", "Contenido con etiqueta de facultad ROSADA arriba a la derecha + subtítulo rosado y texto corto.", "diapositiva_contenido(prs, etiqueta, titulo, color='rosado')"),
    6: ("contenido_etiqueta_morada", "Contenido con etiqueta de facultad MORADA arriba a la derecha + subtítulo morado y texto corto. Es la diapositiva de contenido principal.", "diapositiva_contenido(prs, etiqueta, titulo)"),
    7: ("barra_lateral_morada", "Barra lateral izquierda MORADA de alto completo con el logo blanco; subtítulo y texto a la derecha. Para separar secciones o contenido destacado.", "diapositiva_barra_lateral(prs, titulo, color='morado')"),
    8: ("barra_lateral_rosada", "Barra lateral izquierda ROSADA con el logo blanco; subtítulo rosado y texto a la derecha.", "diapositiva_barra_lateral(prs, titulo, color='rosado')"),
    9: ("foto_con_pie", "Foto centrada con borde blanco y sombra, etiqueta morada encima del borde inferior y un texto corto debajo.", "diapositiva_foto(prs, ruta_imagen, pie, texto)"),
    10: ("lista_de_bloques", "Cuatro bloques apilados a la izquierda: etiqueta de color (morado, rosado, amarillo, cian) + texto corto debajo.", "diapositiva_lista_de_bloques(prs, bloques)"),
    11: ("cierre", "Cierre: foto aérea del campus con banda morada (logo, redes y www.uceva.edu.co) en la imagen de fondo.", "cierre(prs, texto_principal, texto_secundario)"),
}

USO_LAYOUT = {
    "Diapositiva de título": "Diseño usado por casi todas las muestras (1, 2, 4-11) y por la librería. Hereda el fondo del patrón (logo + figuras grises). Sus marcadores (título y subtítulo) NO se usan: se borran y el contenido se pone con formas propias.",
    "En blanco": "Usado por la muestra 3. Mismo fondo del patrón, sin marcadores de título.",
}
USO_LAYOUT_OTRO = "Diseño estándar de Office que trae el archivo; ninguna muestra lo usa. Hereda el fondo del patrón. No hace falta usarlo."

TEXTOS_DE_EJEMPLO = [
    "TÍTULO DE EJEMPLO",
    "Subtítulo de ejemplo",
    "Este es un texto corto de ejemplo",
    "Este es un texto de ejemplo",
    "Facultad de Ciencias Jurídicas y Humanísticas",
]


# --- utilidades -------------------------------------------------------------

def pulg(emu: int) -> float:
    return round(emu / EMU_POR_PULGADA, 3)


def caja(x: int, y: int, cx: int, cy: int) -> dict:
    return {
        "emu": {"x": x, "y": y, "ancho": cx, "alto": cy},
        "pulgadas": {"x": pulg(x), "y": pulg(y), "ancho": pulg(cx), "alto": pulg(cy)},
    }


class Paquete:
    """Acceso de solo lectura a las partes del .pptx (sin descomprimir en disco)."""

    def __init__(self, ruta: Path):
        self.zip = zipfile.ZipFile(ruta)
        self.nombres = set(self.zip.namelist())

    def xml(self, parte: str):
        return etree.fromstring(self.zip.read(parte))

    def bytes(self, parte: str) -> bytes:
        return self.zip.read(parte)

    def rels(self, parte: str) -> dict:
        d, f = posixpath.split(parte)
        ruta_rels = posixpath.join(d, "_rels", f + ".rels")
        if ruta_rels not in self.nombres:
            return {}
        res = {}
        for r in self.xml(ruta_rels).findall("rel:Relationship", NS):
            destino = r.get("Target")
            if r.get("TargetMode") != "External":
                destino = posixpath.normpath(posixpath.join(d, destino))
            res[r.get("Id")] = (r.get("Type").rsplit("/", 1)[-1], destino)
        return res


class Tema:
    def __init__(self, pkg: Paquete, parte_master: str):
        master = pkg.xml(parte_master)
        tema_parte = next(t for tipo, t in pkg.rels(parte_master).values() if tipo == "theme")
        tema = pkg.xml(tema_parte)
        self.colores = {}
        for el in tema.find(".//a:clrScheme", NS):
            hijo = el[0]
            nombre = etree.QName(el).localname
            self.colores[nombre] = (hijo.get("lastClr") or hijo.get("val")).upper()
        mapa = master.find("p:clrMap", NS)
        self.mapa = dict(mapa.attrib)
        self.fuente_menor = tema.find(".//a:minorFont/a:latin", NS).get("typeface")
        self.fuente_mayor = tema.find(".//a:majorFont/a:latin", NS).get("typeface")
        lvl1 = master.find(".//p:otherStyle/a:lvl1pPr/a:defRPr", NS)
        self.tam_por_defecto = int(lvl1.get("sz")) / 100

    def resolver(self, el) -> tuple[str | None, str, float | None]:
        """Devuelve (hex, expresión legible, alfa 0-1) de un srgbClr/schemeClr/sysClr/prstClr."""
        if el is None:
            return None, "", None
        tag = etree.QName(el).localname
        if tag == "srgbClr":
            base, expr = el.get("val").upper(), el.get("val").upper()
        elif tag == "schemeClr":
            nombre = el.get("val")
            real = self.mapa.get(nombre, nombre)
            base, expr = self.colores[real], f"tema:{nombre}"
        elif tag == "sysClr":
            base, expr = el.get("lastClr").upper(), f"sistema:{el.get('val')}"
        elif tag == "prstClr":
            base = {"black": "000000", "white": "FFFFFF"}.get(el.get("val"), "000000")
            expr = f"predefinido:{el.get('val')}"
        else:
            return None, tag, None
        r, g, b = (int(base[i:i + 2], 16) / 255 for i in (0, 2, 4))
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        alfa = None
        for mod in el:
            n, v = etree.QName(mod).localname, int(mod.get("val")) / 100000
            if n == "lumMod":
                l *= v
                expr += f" lumMod {v:.0%}"
            elif n == "lumOff":
                l += v
                expr += f" lumOff {v:.0%}"
            elif n == "alpha":
                alfa = round(v, 4)
        r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), s)
        hexa = "%02X%02X%02X" % tuple(round(c * 255) for c in (r, g, b))
        return hexa, expr, alfa


def color_de(cont, tema: Tema):
    """Lee solidFill/noFill dentro de un spPr, ln o rPr."""
    if cont is None:
        return None
    if cont.find("a:noFill", NS) is not None:
        return {"tipo": "ninguno"}
    sf = cont.find("a:solidFill", NS)
    if sf is not None and len(sf):
        hexa, expr, alfa = tema.resolver(sf[0])
        d = {"tipo": "solido", "hex": hexa, "origen": expr}
        if alfa is not None:
            d["opacidad"] = alfa
        return d
    bf = cont.find("a:blipFill", NS)
    if bf is not None:
        return {"tipo": "imagen"}
    return None


def sombra_de(cont, tema: Tema):
    if cont is None:
        return None
    sh = cont.find("a:effectLst/a:outerShdw", NS)
    if sh is None:
        return None
    hexa, _, alfa = tema.resolver(sh[0] if len(sh) else None)
    return {
        "desenfoque_emu": int(sh.get("blurRad", 0)),
        "distancia_emu": int(sh.get("dist", 0)),
        "direccion_grados": int(sh.get("dir", 0)) / 60000,
        "color": hexa,
        "opacidad": alfa,
        "xml": re.sub(r' xmlns:\w+="[^"]*"', "", etree.tostring(sh, encoding="unicode")),
    }


def color_texto_por_estilo(sp, tema: Tema):
    fr = sp.find("p:style/a:fontRef", NS)
    if fr is not None and len(fr):
        return tema.resolver(fr[0])
    return tema.resolver(etree.fromstring(f'<a:schemeClr xmlns:a="{NS["a"]}" val="tx1"/>'))


def leer_texto(sp, tema: Tema):
    tx = sp.find("p:txBody", NS)
    if tx is None:
        return None, None
    body = tx.find("a:bodyPr", NS)
    ajuste = "ninguno"
    if body.find("a:spAutoFit", NS) is not None:
        ajuste = "forma_se_ajusta_al_texto"
    elif body.find("a:normAutofit", NS) is not None:
        ajuste = "texto_se_encoge"
    margen = {k: int(body.get(k, d)) for k, d in
              (("lIns", 91440), ("tIns", 45720), ("rIns", 91440), ("bIns", 45720))}
    cuerpo = {
        "anclaje": {"ctr": "centro", "b": "abajo", "t": "arriba"}.get(body.get("anchor", "t"), body.get("anchor")),
        "ajuste": ajuste,
        "ajuste_de_linea": body.get("wrap", "square") == "square",
        "margen_interno_emu": margen,
    }
    parrafos = []
    hex_estilo, expr_estilo, _ = color_texto_por_estilo(sp, tema)
    for p in tx.findall("a:p", NS):
        ppr = p.find("a:pPr", NS)
        algn = ppr.get("algn", "l") if ppr is not None else "l"
        runs = []
        for r in p.findall("a:r", NS):
            rpr = r.find("a:rPr", NS)
            rpr = rpr if rpr is not None else etree.Element("x")
            col = color_de(rpr, tema)
            latin = rpr.find("a:latin", NS)
            fuente = latin.get("typeface") if latin is not None else f"{tema.fuente_menor} (heredada del tema)"
            sz = rpr.get("sz")
            run = {
                "texto": r.findtext("a:t", namespaces=NS),
                "fuente": fuente,
                "tam_pt": int(sz) / 100 if sz else tema.tam_por_defecto,
                "tam_heredado": sz is None,
                "negrita": rpr.get("b") == "1",
                "color": col["hex"] if col else hex_estilo,
                "color_origen": col["origen"] if col else f"{expr_estilo} (estilo de la forma)",
            }
            sombra = sombra_de(rpr, tema)
            if sombra:
                run["sombra"] = sombra
            runs.append(run)
        if runs:
            parrafos.append({"alineacion": {"l": "izquierda", "ctr": "centro", "r": "derecha", "just": "justificado"}.get(algn, algn), "runs": runs})
    return cuerpo, parrafos


def leer_formas(pkg: Paquete, parte: str, tema: Tema, nombres_img: dict):
    raiz = pkg.xml(parte)
    rels = pkg.rels(parte)
    formas = []
    arbol = raiz.find("p:cSld/p:spTree", NS)
    for el in arbol:
        tag = etree.QName(el).localname
        if tag not in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp"):
            continue
        cnv = el.find(".//p:cNvPr", NS)
        xfrm = el.find("p:spPr/a:xfrm", NS)
        if xfrm is None:
            xfrm = el.find("p:xfrm", NS)
        forma = {"tipo": {"sp": "forma", "pic": "imagen", "graphicFrame": "tabla_o_grafico",
                          "grpSp": "grupo", "cxnSp": "conector"}[tag],
                 "nombre": cnv.get("name"), "id": int(cnv.get("id"))}
        ph = el.find(".//p:nvPr/p:ph", NS)
        if ph is not None:
            forma["marcador"] = {"tipo": ph.get("type", "body"), "idx": ph.get("idx")}
        if el.find(".//p:cNvSpPr", NS) is not None and el.find(".//p:cNvSpPr", NS).get("txBox") == "1":
            forma["tipo"] = "cuadro_de_texto"
        if xfrm is not None and xfrm.find("a:off", NS) is not None:
            off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
            forma["posicion"] = caja(int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
        spPr = el.find("p:spPr", NS)
        geom = spPr.find("a:prstGeom", NS) if spPr is not None else None
        if geom is not None:
            forma["geometria"] = {"rect": "rectangulo", "roundRect": "rectangulo_redondeado"}.get(geom.get("prst"), geom.get("prst"))
            gd = geom.find("a:avLst/a:gd", NS)
            if geom.get("prst") == "roundRect":
                forma["radio_esquinas"] = gd.get("fmla") if gd is not None else "por defecto (adj 16667 = 16,7 % del lado corto)"
        relleno = color_de(spPr, tema)
        if relleno is None and el.find("p:style/a:fillRef", NS) is not None and tag == "sp":
            fr = el.find("p:style/a:fillRef", NS)
            if fr.get("idx") != "0":
                hexa, expr, _ = tema.resolver(fr[0])
                relleno = {"tipo": "solido", "hex": hexa, "origen": expr + " (estilo de la forma)"}
        if relleno:
            forma["relleno"] = relleno
        ln = spPr.find("a:ln", NS) if spPr is not None else None
        if ln is not None:
            linea = color_de(ln, tema) or {}
            if ln.get("w"):
                linea["ancho_pt"] = int(ln.get("w")) / EMU_POR_PT
            dash = ln.find("a:prstDash", NS)
            if dash is not None:
                linea["guion"] = {"dash": "discontinua", "sysDash": "punteada"}.get(dash.get("val"), dash.get("val"))
            forma["linea"] = linea
        sombra = sombra_de(spPr, tema)
        if sombra:
            forma["sombra"] = sombra
        if tag == "pic":
            blip = el.find(".//a:blip", NS)
            destino = rels[blip.get(f"{{{NS['r']}}}embed")][1]
            forma["imagen"] = nombres_img.get(destino, destino)
        cuerpo, parrafos = leer_texto(el, tema)
        if parrafos is not None:
            forma["cuadro_texto"] = cuerpo
            forma["parrafos"] = parrafos
        formas.append(forma)
    return formas


def fondo_de(pkg: Paquete, parte: str, nombres_img: dict):
    raiz = pkg.xml(parte)
    blip = raiz.find("p:cSld/p:bg/p:bgPr/a:blipFill/a:blip", NS)
    if blip is None:
        return None
    destino = pkg.rels(parte)[blip.get(f"{{{NS['r']}}}embed")][1]
    return nombres_img.get(destino, destino)


def bbox_color(img: Image.Image, region) -> tuple[int, int, int, int]:
    """Caja (px) que encierra los píxeles con color (saturados) dentro de la región (fracciones)."""
    W, H = img.size
    x0, y0, x1, y1 = int(region[0] * W), int(region[1] * H), int(region[2] * W), int(region[3] * H)
    _, s, v = img.convert("RGB").convert("HSV").crop((x0, y0, x1, y1)).split()
    # saturación > 35 % y brillo > 20 %: descarta el blanco y los grises del fondo
    m = ImageChops.multiply(s.point(lambda q: 255 if q > 89 else 0), v.point(lambda q: 255 if q > 51 else 0))
    bb = m.getbbox()
    return (bb[0] + x0, bb[1] + y0, bb[2] + x0, bb[3] + y0)


def bandas_moradas(img: Image.Image) -> tuple[int, int]:
    """Filas (px) donde el morado de la banda del cierre ocupa la mayor parte del ancho."""
    W, H = img.size
    px = img.convert("RGB").load()
    filas = [y for y in range(H)
             if sum(1 for x in range(0, W, 10)
                    if abs(px[x, y][0] - 0x60) < 40 and px[x, y][1] < 60 and px[x, y][2] > 120) > (W / 10) * 0.6]
    return min(filas), max(filas) + 1


def px_a_emu(px: int, total_px: int, total_emu: int) -> int:
    return round(px * total_emu / total_px)


# --- extracción principal ---------------------------------------------------

def extraer(pptx: Path, salida: Path, previews: bool = True) -> None:
    pkg = Paquete(pptx)
    pres = pkg.xml("ppt/presentation.xml")
    pres_rels = pkg.rels("ppt/presentation.xml")
    sz = pres.find("p:sldSz", NS)
    ANCHO, ALTO = int(sz.get("cx")), int(sz.get("cy"))

    master_rid = pres.find("p:sldMasterIdLst/p:sldMasterId", NS).get(f"{{{NS['r']}}}id")
    parte_master = pres_rels[master_rid][1]
    tema = Tema(pkg, parte_master)

    slides = []
    for i, s in enumerate(pres.findall("p:sldIdLst/p:sldId", NS), start=1):
        slides.append((i, int(s.get("id")), pres_rels[s.get(f"{{{NS['r']}}}id")][1]))

    # 1) Imágenes: nombre claro según dónde se usan.
    nombres_img: dict[str, str] = {}
    def registrar(destino, rol):
        if destino in nombres_img:
            return
        ext = Path(destino).suffix.lower()
        base = ROL_IMAGEN.get(rol, f"imagen_{Path(destino).stem}")
        nombres_img[destino] = f"assets/img/{base}{ext}"

    master = pkg.xml(parte_master)
    mblip = master.find("p:cSld/p:bg/p:bgPr/a:blipFill/a:blip", NS)
    registrar(pkg.rels(parte_master)[mblip.get(f"{{{NS['r']}}}embed")][1], "fondo_patron")
    for num, _, parte in slides:
        rels = pkg.rels(parte)
        raiz = pkg.xml(parte)
        b = raiz.find("p:cSld/p:bg/p:bgPr/a:blipFill/a:blip", NS)
        if b is not None:
            registrar(rels[b.get(f"{{{NS['r']}}}embed")][1], f"fondo_slide{num}")
        for pic in raiz.iter(f"{{{NS['p']}}}pic"):
            nombre = pic.find(".//p:cNvPr", NS).get("name")
            bl = pic.find(".//a:blip", NS)
            registrar(rels[bl.get(f"{{{NS['r']}}}embed")][1], f"imagen_slide{num}_{nombre}")

    dir_img = salida / "assets" / "img"
    if dir_img.exists():
        shutil.rmtree(dir_img)
    dir_img.mkdir(parents=True)
    hashes = {}
    for destino, rel_path in sorted(nombres_img.items(), key=lambda kv: kv[1]):
        datos = pkg.bytes(destino)
        (salida / rel_path).write_bytes(datos)
        hashes[rel_path] = hashlib.sha1(datos).hexdigest()
    otros = sorted(n for n in pkg.nombres if n.startswith("ppt/media/") and n not in nombres_img)
    for n in otros:
        # Formatos vectoriales (svg/emf/wmf): se copian y, si Pillow puede, se pasan a PNG.
        nombres_img[n] = f"assets/img/{Path(n).name}"
        (salida / nombres_img[n]).write_bytes(pkg.bytes(n))
    for n in otros:
        if Path(n).suffix.lower() in (".emf", ".wmf", ".svg"):
            try:
                Image.open(io.BytesIO(pkg.bytes(n))).save(dir_img / (Path(n).stem + ".png"))
            except Exception:
                print(f"  aviso: no pude convertir {n} a PNG (falta un conversor vectorial)")

    # Recorte del logo a color desde el fondo de la portada (mayor resolución disponible).
    nombre_portada = next(v for v in nombres_img.values() if "fondo_portada" in v)
    nombre_contenido = next(v for v in nombres_img.values() if "fondo_contenido" in v)
    img_portada = Image.open(salida / nombre_portada)
    bb = bbox_color(img_portada, (0, 0, 0.4, 0.2))
    pad = 12
    logo = img_portada.convert("RGB").crop((bb[0] - pad, bb[1] - pad, bb[2] + pad, bb[3] + pad))
    logo.save(dir_img / "logo_55_uceva_color_recorte.png", optimize=False)
    hashes["assets/img/logo_55_uceva_color_recorte.png"] = hashlib.sha1((dir_img / "logo_55_uceva_color_recorte.png").read_bytes()).hexdigest()

    # 2) Zonas protegidas medidas en las imágenes de fondo.
    def zona_logo(ruta_img, region):
        img = Image.open(salida / ruta_img)
        W, H = img.size
        b = bbox_color(img, region)
        return {"x": 0, "y": 0,
                "ancho": px_a_emu(b[2], W, ANCHO) + 91440,
                "alto": px_a_emu(b[3], H, ALTO) + 91440,
                "logo_medido_emu": {"x": px_a_emu(b[0], W, ANCHO), "y": px_a_emu(b[1], H, ALTO),
                                    "ancho": px_a_emu(b[2] - b[0], W, ANCHO), "alto": px_a_emu(b[3] - b[1], H, ALTO)}}

    z_contenido = zona_logo(nombre_contenido, (0, 0, 0.4, 0.25))
    z_portada = zona_logo(nombre_portada, (0, 0, 0.4, 0.2))
    nombre_cierre = next(v for v in nombres_img.values() if "fondo_cierre" in v)
    img_cierre = Image.open(salida / nombre_cierre)
    f0, f1 = bandas_moradas(img_cierre)
    banda_cierre = {"x": 0, "y": px_a_emu(f0, img_cierre.size[1], ALTO), "ancho": ANCHO,
                    "alto": px_a_emu(f1 - f0, img_cierre.size[1], ALTO)}
    # "Bienvenidos" pequeño de la portada (abajo a la izquierda): caja de color en la mitad inferior izquierda.
    bb_b = bbox_color(img_portada, (0, 0.45, 0.5, 1.0))
    W, H = img_portada.size
    arte_portada = {"x": px_a_emu(bb_b[0], W, ANCHO), "y": px_a_emu(bb_b[1], H, ALTO),
                    "ancho": px_a_emu(bb_b[2] - bb_b[0], W, ANCHO), "alto": px_a_emu(bb_b[3] - bb_b[1], H, ALTO)}

    # 3) slides.json
    datos_slides = []
    nombres_layout = {}
    usos_layout: dict[str, list[int]] = {}
    colores_vistos: dict[str, set] = {}
    tamanos_vistos: set = set()
    fuentes_vistas: set = set()
    for num, sid, parte in slides:
        rels = pkg.rels(parte)
        lay = next(t for tipo, t in rels.values() if tipo == "slideLayout")
        if lay not in nombres_layout:
            nombres_layout[lay] = pkg.xml(lay).find("p:cSld", NS).get("name")
        usos_layout.setdefault(lay, []).append(num)
        formas = leer_formas(pkg, parte, tema, nombres_img)
        fondo = fondo_de(pkg, parte, nombres_img)
        clave, desc, funcion = DESCRIPCION_SLIDE.get(num, (f"muestra_{num}", "", ""))
        for f in formas:
            for k in ("relleno", "linea"):
                if f.get(k, {}).get("hex"):
                    colores_vistos.setdefault(f[k]["hex"], set()).add(num)
            for p in f.get("parrafos", []):
                for r in p["runs"]:
                    colores_vistos.setdefault(r["color"], set()).add(num)
                    tamanos_vistos.add(r["tam_pt"])
                    fuentes_vistas.add(r["fuente"])
                    if r.get("sombra"):
                        colores_vistos.setdefault(r["sombra"]["color"], set()).add(num)
        datos_slides.append({
            "numero": num,
            "id_interno": sid,
            "clave": clave,
            "cuando_usarla": desc,
            "funcion_lib": funcion,
            "layout": nombres_layout[lay],
            "fondo": {"tipo": "propio", "imagen": fondo} if fondo else {"tipo": "heredado_del_patron", "imagen": nombre_contenido},
            "vista_previa": f"assets/previews/slide{num:02d}.png",
            "textos_de_ejemplo": [r["texto"] for f in formas for p in f.get("parrafos", []) for r in p["runs"]],
            "formas": formas,
        })

    # 4) layouts.json
    datos_layouts = []
    m_rels = pkg.rels(parte_master)
    master_ph = {}
    for sp in master.find("p:cSld/p:spTree", NS).findall("p:sp", NS):
        ph = sp.find(".//p:nvPr/p:ph", NS)
        off, ext = sp.find("p:spPr/a:xfrm/a:off", NS), sp.find("p:spPr/a:xfrm/a:ext", NS)
        if ph is not None and off is not None:
            master_ph[ph.get("type", "body")] = (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
    orden_layouts = [m_rels[e.get(f"{{{NS['r']}}}id")][1] for e in master.findall("p:sldLayoutIdLst/p:sldLayoutId", NS)]
    for idx, lay in enumerate(orden_layouts):
        raiz = pkg.xml(lay)
        nombre = raiz.find("p:cSld", NS).get("name")
        phs = []
        for sp in raiz.find("p:cSld/p:spTree", NS).findall("p:sp", NS):
            ph = sp.find(".//p:nvPr/p:ph", NS)
            if ph is None:
                continue
            tipo = ph.get("type", "body")
            off, ext = sp.find("p:spPr/a:xfrm/a:off", NS), sp.find("p:spPr/a:xfrm/a:ext", NS)
            if off is not None:
                pos, origen = (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))), "layout"
            else:
                clave_m = {"ctrTitle": "title", "subTitle": "body"}.get(tipo, tipo)
                pos, origen = master_ph.get(clave_m), "heredada_del_patron"
            d = {"tipo": tipo, "idx": ph.get("idx"), "nombre": sp.find(".//p:cNvPr", NS).get("name"),
                 "posicion_origen": origen}
            if pos:
                d["posicion"] = caja(*pos)
            phs.append(d)
        datos_layouts.append({
            "indice": idx,
            "nombre": nombre,
            "tipo": raiz.get("type"),
            "parte": lay,
            "usado_por_muestras": usos_layout.get(lay, []),
            "para_que_sirve": USO_LAYOUT.get(nombre, USO_LAYOUT_OTRO),
            "fondo": "heredado del patrón: " + nombre_contenido,
            "marcadores": phs,
        })

    # 5) tokens.json
    gris_texto = None
    for s in datos_slides:
        for f in s["formas"]:
            for p in f.get("parrafos", []):
                for r in p["runs"]:
                    if r["color_origen"].startswith("tema:bg2"):
                        gris_texto = (r["color"], r["color_origen"])
    gris_claro = tema.colores[tema.mapa["bg2"]]
    paleta = {}
    for hexa in sorted(colores_vistos, key=lambda h: (h not in NOMBRES_COLOR, h)):
        nombre = NOMBRES_COLOR.get(hexa)
        if gris_texto and hexa == gris_texto[0]:
            nombre = "gris_texto"
        if not nombre:
            nombre = f"color_{hexa.lower()}"
            print(f"  aviso: color {hexa} sin nombre en NOMBRES_COLOR")
        paleta[nombre] = {"hex": hexa, "rgb": [int(hexa[i:i + 2], 16) for i in (0, 2, 4)],
                          "uso": USO_COLOR.get(nombre, ""), "visto_en_muestras": sorted(colores_vistos[hexa])}
        if nombre == "gris_texto":
            paleta[nombre]["origen"] = f"{gris_texto[1]} (Fondo 2 del tema = {gris_claro}, oscurecido al 25 %)"
    paleta["gris_claro"] = {"hex": gris_claro, "rgb": [int(gris_claro[i:i + 2], 16) for i in (0, 2, 4)],
                            "uso": USO_COLOR["gris_claro"], "visto_en_muestras": [], "origen": "tema: lt2 / bg2"}

    def mezclar(hexa, factor, con="000000"):
        a = [int(hexa[i:i + 2], 16) for i in (0, 2, 4)]
        b = [int(con[i:i + 2], 16) for i in (0, 2, 4)]
        return "%02X%02X%02X" % tuple(round(x * (1 - factor) + y * factor) for x, y in zip(a, b))

    morado = paleta["morado"]["hex"]
    derivados = {
        "morado_oscuro": {"hex": mezclar(morado, 0.70), "uso": "Fondo de bloques de código (texto blanco encima).",
                          "origen": f"morado {morado} mezclado 70 % con negro"},
        "morado_suave": {"hex": mezclar(morado, 0.92, "FFFFFF"), "uso": "Fondo de filas alternas de tablas y de tarjetas suaves.",
                         "origen": f"morado {morado} mezclado 92 % con blanco"},
    }

    etiqueta = next(f for f in datos_slides[5]["formas"] if f.get("geometria") == "rectangulo_redondeado")
    titulo_cont = next(f for f in datos_slides[5]["formas"] if f["tipo"] == "cuadro_de_texto")
    titulo_portada = next(f for f in datos_slides[1]["formas"] if f.get("relleno", {}).get("hex") == morado)
    marco_portada = next(f for f in datos_slides[1]["formas"] if f.get("linea", {}).get("guion"))
    sub_portada = next(f for f in datos_slides[1]["formas"] if f["tipo"] == "cuadro_de_texto")
    barra = next(f for f in datos_slides[6]["formas"] if f.get("geometria") == "rectangulo" and f.get("relleno", {}).get("hex") == morado)
    logo_barra = next(f for f in datos_slides[6]["formas"] if f["tipo"] == "imagen")
    texto_barra = next(f for f in datos_slides[6]["formas"] if f["tipo"] == "cuadro_de_texto")
    foto = next(f for f in datos_slides[8]["formas"] if f["tipo"] == "imagen")
    pie_etq = next(f for f in datos_slides[8]["formas"] if f.get("geometria") == "rectangulo_redondeado")
    pie_txt = next(f for f in datos_slides[8]["formas"] if f["tipo"] == "cuadro_de_texto")
    bloques = [f for f in datos_slides[9]["formas"]]

    margen_lateral = 365760  # 0,4"
    tokens = {
        "_leeme": "Generado por scripts/extraer_plantilla.py desde el .pptx. No editar a mano: vuelve a ejecutar el script.",
        "hoja": {"ancho_emu": ANCHO, "alto_emu": ALTO, "ancho_pulgadas": pulg(ANCHO), "alto_pulgadas": pulg(ALTO),
                 "proporcion": "16:9", "emu_por_pulgada": EMU_POR_PULGADA, "emu_por_punto": EMU_POR_PT},
        "colores": paleta,
        "colores_derivados": derivados,
        "fuentes": {
            "texto": "Arial",
            "titulo_portada": "Arial Black",
            "codigo": "Courier New",
            "vistas_en_la_plantilla": sorted(fuentes_vistas),
            "fuente_del_tema": f"{tema.fuente_menor} / {tema.fuente_mayor}",
            "advertencia": f"Si un texto no declara fuente, hereda {tema.fuente_menor} del tema. Siempre declara Arial.",
        },
        "tamanos_pt": {
            "vistos_en_la_plantilla": sorted(tamanos_vistos),
            "por_defecto_heredado": tema.tam_por_defecto,
            "titulo_portada": 32, "titulo_portada_minimo": 24,
            "titulo_diapositiva": 24,
            "subtitulo": 18,
            "cuerpo": 16, "cuerpo_minimo": 14,
            "etiqueta": 16, "etiqueta_bloque": 18,
            "pie_minimo": 12,
            "codigo": 14, "codigo_minimo": 12,
            "numero_grande": 54,
            "nota": "Los mínimos son reglas de la librería (no vienen del archivo). El texto corto de las muestras 4-8 no declara tamaño y hereda 18 pt; su marca de fin de párrafo dice 16 pt.",
        },
        "margenes": {
            "lateral_minimo_emu": margen_lateral, "lateral_minimo_pulgadas": pulg(margen_lateral),
            "inferior_minimo_emu": margen_lateral, "inferior_minimo_pulgadas": pulg(margen_lateral),
            "x_texto_plantilla_emu": titulo_cont["posicion"]["emu"]["x"],
            "nota": "La plantilla alinea el texto en x=%d EMU (%.2f\"). La librería usa ese valor como margen izquierdo." % (
                titulo_cont["posicion"]["emu"]["x"], pulg(titulo_cont["posicion"]["emu"]["x"])),
        },
        "zonas_protegidas": {
            "logo_contenido": {**z_contenido, "aplica_a": "diapositivas con el fondo del patrón (contenido, en blanco, foto, bloques)"},
            "logo_portada": {**z_portada, "aplica_a": "portada y bienvenida"},
            "etiqueta_facultad": {**etiqueta["posicion"]["emu"], "aplica_a": "diapositivas de contenido con etiqueta"},
            "barra_lateral": {**barra["posicion"]["emu"], "aplica_a": "diapositivas con barra lateral"},
            "banda_cierre": {**banda_cierre, "aplica_a": "cierre (banda morada con logo y web, dentro de la imagen)"},
            "arte_bienvenidos_portada": {**arte_portada, "aplica_a": "portada (arte 'Bienvenidos' pequeño abajo a la izquierda)"},
        },
        "elementos": {
            "etiqueta_facultad": {"posicion": etiqueta["posicion"], "geometria": "rectangulo_redondeado",
                                  "texto": {k: etiqueta["parrafos"][0]["runs"][0][k] for k in ("fuente", "tam_pt", "negrita", "color")},
                                  "sombra_texto_xml": etiqueta["parrafos"][0]["runs"][0]["sombra"]["xml"],
                                  "colores_en_muestras": {"4": "amarillo", "5": "rosado", "6": "morado"}},
            "titulo_contenido": {"posicion": titulo_cont["posicion"],
                                 "texto": {k: titulo_cont["parrafos"][0]["runs"][0][k] for k in ("fuente", "tam_pt", "negrita", "color")},
                                 "texto_corto": {k: titulo_cont["parrafos"][1]["runs"][0][k] for k in ("fuente", "tam_pt", "negrita", "color", "tam_heredado")}},
            "titulo_portada": {"posicion": titulo_portada["posicion"],
                               "texto": {k: titulo_portada["parrafos"][0]["runs"][0][k] for k in ("fuente", "tam_pt", "negrita", "color")}},
            "marco_portada": {"posicion": marco_portada["posicion"], "linea": marco_portada["linea"],
                              "nota": "Se sale 0,2\" por la derecha de la hoja a propósito (sangrado)."},
            "subtitulo_portada": {"posicion": sub_portada["posicion"],
                                  "texto": {k: sub_portada["parrafos"][0]["runs"][0][k] for k in ("fuente", "tam_pt", "negrita", "color")}},
            "barra_lateral": {"posicion": barra["posicion"], "colores_en_muestras": {"7": "morado", "8": "rosado"}},
            "logo_blanco_barra": {"posicion": logo_barra["posicion"], "imagen": logo_barra["imagen"]},
            "texto_con_barra": {"posicion": texto_barra["posicion"]},
            "foto": {"posicion": foto["posicion"], "linea": foto.get("linea"), "sombra_xml": foto["sombra"]["xml"]},
            "pie_foto_etiqueta": {"posicion": pie_etq["posicion"],
                                  "texto": {k: pie_etq["parrafos"][0]["runs"][0][k] for k in ("fuente", "tam_pt", "negrita", "color")}},
            "pie_foto_texto": {"posicion": pie_txt["posicion"],
                               "texto": {k: pie_txt["parrafos"][0]["runs"][0][k] for k in ("fuente", "tam_pt", "negrita", "color")}},
            "bloques": [{"posicion": b["posicion"]["emu"], "relleno": b.get("relleno", {}).get("hex"),
                         "texto": b["parrafos"][0]["runs"][0]["texto"] if b.get("parrafos") else None} for b in bloques],
        },
        "imagenes": {Path(p).stem: {"ruta": p, "sha1": h} for p, h in sorted(hashes.items())},
        "textos_de_ejemplo": TEXTOS_DE_EJEMPLO,
    }

    dir_spec = salida / "spec"
    dir_spec.mkdir(parents=True, exist_ok=True)
    def escribir(nombre, datos):
        (dir_spec / nombre).write_text(json.dumps(datos, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    escribir("tokens.json", tokens)
    escribir("layouts.json", {"_leeme": tokens["_leeme"], "patron": {"parte": parte_master, "fondo": nombre_contenido,
                                                                     "marcadores": {k: caja(*v) for k, v in master_ph.items()}},
                              "layouts": datos_layouts})
    escribir("slides.json", {"_leeme": tokens["_leeme"], "hoja": tokens["hoja"], "slides": datos_slides})

    verificar(pptx, datos_slides)

    if previews:
        dir_prev = salida / "assets" / "previews"
        r = subprocess.run(["bash", str(RAIZ / "scripts" / "render_previews.sh"), str(pptx), str(dir_prev)],
                           capture_output=True, text=True)
        print(("  " + r.stdout.strip()) if r.returncode == 0 else f"  aviso: no se generaron vistas previas: {r.stderr.strip()}")

    print(f"Listo: {len(datos_slides)} diapositivas, {len(datos_layouts)} layouts, "
          f"{len(paleta)} colores, {len(hashes)} imágenes → {salida}")


def verificar(pptx: Path, datos_slides: list) -> None:
    """Vuelve a leer el .pptx con python-pptx y compara posiciones y colores con lo extraído."""
    from pptx import Presentation
    prs = Presentation(str(pptx))
    revisadas = 0
    for s, d in zip(prs.slides, datos_slides):
        por_id = {sh.shape_id: sh for sh in s.shapes}
        for f in d["formas"]:
            sh = por_id[f["id"]]
            e = f["posicion"]["emu"]
            assert (sh.left, sh.top, sh.width, sh.height) == (e["x"], e["y"], e["ancho"], e["alto"]), (d["numero"], f["nombre"])
            if f.get("relleno", {}).get("tipo") == "solido" and "estilo" not in f["relleno"]["origen"] and not f["relleno"]["origen"].startswith("tema"):
                assert str(sh.fill.fore_color.rgb) == f["relleno"]["hex"], (d["numero"], f["nombre"])
            revisadas += 1
    print(f"  verificación: {revisadas} formas releídas con python-pptx coinciden en posición y color")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pptx", type=Path, default=PPTX_POR_DEFECTO)
    ap.add_argument("--salida", type=Path, default=RAIZ)
    ap.add_argument("--sin-previews", action="store_true")
    a = ap.parse_args()
    if not a.pptx.exists():
        sys.exit(f"No encuentro la plantilla: {a.pptx}")
    extraer(a.pptx, a.salida, previews=not a.sin_previews)


if __name__ == "__main__":
    main()
